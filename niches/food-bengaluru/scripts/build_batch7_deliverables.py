#!/usr/bin/env python3
"""Render the source-checked Batch 7 Food/Bengaluru deliverables.

This is an incremental pack updater. It appends the six evidence-qualified Batch 7
leads to the existing pack and does not rebuild or overwrite earlier batches. It
refuses to proceed if Batch 7 artifacts or registry entries already exist.
Run from the repository root with the ignored .venv active:
    .venv/bin/python niches/food-bengaluru/scripts/build_batch7_deliverables.py
"""
from __future__ import annotations

import csv
import html
import json
import re
import shutil
import unicodedata
from collections import Counter
from copy import copy
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
DATA_FILE = ROOT / "data" / "batch7_records.json"
WORKBOOK = ROOT / "Food-Bangalore-ALL-IN-ONE.xlsx"
EMAIL_WORKBOOK = ROOT / "Food-Bangalore-EMAILS-ONE-SHEET.xlsx"
BATCH7_EMAIL_WORKBOOK = ROOT / "Food-Bangalore-BATCH-7-EMAILS.xlsx"
EMAIL_CSV = ROOT / "Food-Bangalore-EMAILS-ONE-SHEET.csv"
LEADS_CSV = ROOT / "LEADS-INDEX.csv"
VERIFY_MD = ROOT / "VERIFY-THE-DATA.md"
README_MD = ROOT / "README.md"
ALL_EMAILS_MD = ROOT / "ALL-EMAIL-SEQUENCES.md"
REGISTRY = REPO / "data" / "seen_leads.json"

STEP0_MARKER = "### Step 0 — repository inventory and skill dispositions for Batch 7"
STEP0_FALLBACK = (
    "### Step 0 — repository inventory and skill dispositions for Batch 7\n\n"
    "- Used this repository's REUSABLE-PROMPT.md, the existing Food/Bengaluru archive and the incremental Batch 6 builder as pack/schema references; the Batch 7 builder is an adaptation, not a rebuild.\n"
    "- Used the repository's performance-lead-audit framing, ads/analytics/attribution guardrails and cold-email references only for source-led framing, measurement boundaries and concise Day 1/3/7/14 copy.\n"
    "- External research, audit, analytics/attribution and prospecting references used are recorded in the prior Step 0 inventory. Mock/generated lead pools, Maps-scraping scripts and unsupported projection examples were not run or copied.\n"
    "- Repositories without a relevant usable skill and unrelated repositories are explicitly listed as not used/irrelevant in the preceding Step 0 inventory.\n"
)

NAVY = "0B1B2B"
NAVY_2 = "10263A"
TEAL = "0E7C7B"
GOLD = "F0A500"
PANEL = "F4F7FA"
RULE = "E8EDF3"
MUTED = "64748B"
INK = "172533"
WHITE = "FFFFFF"

LEGACY_BATCH1 = {
    "Licious", "Cothas Coffee", "Anand Sweets", "Chai Point",
    "The Baker's Dozen", "Early Foods", "iD Fresh Food",
    "Third Wave Coffee", "Akshayakalpa Organic", "Milky Mist", "Frozen Bottle",
}
LEGACY_BATCH2 = {
    "SMOOR", "Sid's Farm", "Barbeque Nation", "Millet Amma", "Adukale",
    "Organic Mandya", "Pure & Sure", "Eat Better Co", "Brik Oven", "Araku Coffee",
}
BATCH3_LEADS = {
    "Aubree", "Nandhana Palace", "Byg Brewski", "Hatti Kaapi", "MTR Foods",
    "Popeyes India", "Sagar Ratna", "McDonald's India", "Starbucks India", "Krispy Kreme India",
}
BATCH4_LEADS = {
    "Theobroma", "The Belgian Waffle Co.", "Behrouz Biryani", "Bakingo", "Taco Bell India",
    "Wow! Momo", "Polar Bear", "Jumboking", "Swiggy", "Chaayos",
}
BATCH5_LEADS = {
    "NATUF", "Swish Now", "Thalairaj Biryani", "Tuk Tuk Thai India", "Chelvies Coffee",
    "Liliyum Patisserie Cafe", "Fish Mart", "Mixnosh Art Cafe", "The Cuisine Story", "Sawadee",
}
BATCH6_LEADS = {
    "NIKAA Briyani", "Black Pearl Barbeque", "Suvaii", "Gold Coins Club & Resort",
    "Paint the Town Restaurant", "Chutney Chang", "Xin by MindEscapes",
    "Roast Aroma", "AN’s Events & Caterers",
}
BATCH7_LEADS = {
    "GO DESi", "Boba Bhaii", "Deliciae", "United Telugu Kitchens",
    "The Kind Roastery and Brew Room", "La Pino'z Pizza",
}


def norm_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def norm_domain(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    if "://" not in raw:
        raw = "https://" + raw
    host = (urlparse(raw).hostname or "").casefold()
    return host[4:] if host.startswith("www.") else host


def vintage_for(brand: str) -> str:
    key = norm_name(brand)
    groups = [
        (LEGACY_BATCH1, "LEGACY — Batch 1; prior snapshot, not revalidated on 06 OCT 2026"),
        (LEGACY_BATCH2, "LEGACY — Batch 2; prior snapshot, not revalidated on 06 OCT 2026"),
        (BATCH3_LEADS, "BATCH 3 — source-checked 04 OCT 2026"),
        (BATCH4_LEADS, "BATCH 4 — source-checked 04 OCT 2026"),
        (BATCH5_LEADS, "BATCH 5 — source-checked 04 OCT 2026"),
        (BATCH6_LEADS, "BATCH 6 — source-checked 05 OCT 2026"),
        (BATCH7_LEADS, "BATCH 7 — source-checked 06 OCT 2026"),
    ]
    for group, label in groups:
        if any(norm_name(x) == key for x in group):
            return label
    return "Unclassified legacy — batch/data vintage not mapped"

def load_data():
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Missing source file: {DATA_FILE}")
    data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
    leads = data.get("leads", [])
    assert data.get("batch") == 7, "Expected batch 7 input"
    working_target = int(data.get("working_target", 10))
    selected_count = int(data.get("selected_count", len(leads)))
    assert working_target == 10, f"Unexpected working target: {working_target}"
    assert 1 <= len(leads) <= working_target, f"Expected 1–{working_target} evidence-qualified leads; found {len(leads)}"
    assert selected_count == len(leads), f"Selected-count metadata mismatch: {selected_count} vs {len(leads)}"
    assert len({x["slug"] for x in leads}) == len(leads), "Duplicate slugs in batch input"
    assert len({norm_name(x["brand"]) for x in leads}) == len(leads), "Duplicate brands in batch input"
    dimension_names = data.get("score_dimensions", [])
    assert len(dimension_names) == 5, "Score definition must have five dimensions"
    for lead in leads:
        assert lead.get("contact_source"), f"Missing contact provenance: {lead['brand']}"
        assert len(lead.get("email_notes", [])) == 4 and len(lead.get("email_subjects", [])) == 4, f"Expected four outreach emails: {lead['brand']}"
        assert len(lead.get("findings", [])) >= 1, f"No findings: {lead['brand']}"
        assert len(lead.get("score_values", [])) == 5 and len(lead.get("score_bases", [])) == 5, f"Score must have five dimensions: {lead['brand']}"
        assert sum(int(x) for x in lead["score_values"]) <= 20, f"Invalid score total: {lead['brand']}"
        source_ids = {s["id"] for s in lead.get("sources_detail", [])}
        assert source_ids and len(source_ids) == len(lead["sources_detail"]), f"Missing/duplicate source IDs: {lead['brand']}"
        assert "S1" in source_ids and "S2" in source_ids, f"Missing ad/contact source: {lead['brand']}"
        email_source_ids = lead.get("email_source_ids", [])
        assert len(email_source_ids) == 4, f"Each email needs source IDs: {lead['brand']}"
        assert all(ids and set(ids).issubset(source_ids) for ids in email_source_ids), f"Unresolved email source IDs: {lead['brand']}"
        for finding in lead.get("findings", []):
            assert set(finding.get("source_ids", [])).issubset(source_ids), f"Unresolved finding source: {lead['brand']}"
        for sid in lead.get("contact_source_detail", {}).get("source_ids", []):
            assert sid in source_ids, f"Unresolved contact provenance source: {lead['brand']}"
        email = lead.get("contact", {}).get("email", "").casefold()
        source_text = " ".join([lead.get("contact_source", "")] + [s.get("observed", "") + " " + s.get("url", "") for s in lead["sources_detail"]]).casefold()
        assert email and email in source_text, f"Published email not present in cited first-party provenance: {lead['brand']}"

        meta = lead.pop("meta")
        google = lead["google_evidence"]
        assert meta.get("status") == "Active" and meta.get("sponsored") is True, f"Selected ad is not verified Active/Sponsored: {lead['brand']}"
        assert meta.get("checked_on") == "2026-10-06", f"Ad capture date mismatch: {lead['brand']}"
        assert str(meta.get("library_id", "")) in meta.get("direct_url", "") and str(meta.get("page_id", "")) in meta.get("direct_url", ""), f"Direct Library URL missing IDs: {lead['brand']}"
        direct_url = meta.get("direct_url") or f"https://www.facebook.com/ads/library/?active_status=active&ad_type=all&country=IN&id={meta['library_id']}&view_all_page_id={meta['page_id']}"
        ad = {
            "platform": "Meta Ad Library",
            "status": f"Active and Sponsored on {meta['checked_on']}",
            "active": True, "sponsored": True, "checked_on": meta["checked_on"],
            "library_id": str(meta["library_id"]), "page_id": str(meta["page_id"]),
            "page_name": meta["page_name"], "page_profile_url": meta.get("page_profile_url", ""),
            "started_on": meta["started_on"],
            "creative_excerpt": meta["creative_excerpt"], "cta": meta["cta"],
            "destination_url": meta.get("destination_url"),
            "destination_status": meta.get("destination_status", "Destination not exposed; not tested."),
            "source_id": meta.get("source_id", "S1"), "direct_url": direct_url,
        }
        lead["ad_evidence"] = [ad]
        lead["batch"] = 7
        lead["verified_on"] = "06 OCT 2026"
        values = [int(x) for x in lead.pop("score_values")]
        bases = lead.pop("score_bases")
        total = sum(values)
        lead["score_20"] = total
        lead["score"] = total / 2
        lead["score_band"] = "High public-signal fit" if total >= 18 else "Good public-signal fit" if total >= 15 else "Moderate public-signal fit"
        lead["score_note"] = data["score_note"]
        lead["score_breakdown"] = [{"dimension": name, "score": value, "basis": basis} for name, value, basis in zip(dimension_names, values, bases)]
        gstatus = google["status"].casefold()
        if "no domain-query results" in gstatus or "zero domain-query results" in gstatus:
            lead["ads_active"] = "Google Ads Transparency — " + google["status"]
        else:
            lead["ads_active"] = "Google Ads Transparency — " + google["status"]
        lead["sources"] = [f"{s['id']} — {s['label']}; checked {s.get('verified_on','2026-10-06')}; {s.get('observed','')}" for s in lead["sources_detail"]]
        lead["coverage"] = [
            {"surface":"Meta paid activity","status":"MEASURED — selected direct card active on capture date","method":f"Direct Library ID {ad['library_id']} and page ID {ad['page_id']}; advertiser label, Sponsored/Active status and start date checked. One selected card, not an account-wide count.","date":"06 OCT 2026"},
            {"surface":"Google Ads Transparency","status":google["status"],"method":google["detail"],"date":"06 OCT 2026"},
            {"surface":"Bengaluru context and first-party contact","status":lead.get("local_evidence_status", "MEASURED — public first-party page content"),"method":lead["market_signal"] + " " + lead["contact"]["route_role"],"date":"06 OCT 2026"},
            {"surface":"CTA / destination path","status":"MEASURED — public card label/link only" if ad.get("destination_url") else "NOT MEASURED — destination not exposed in captured card","method":ad["destination_status"],"date":"06 OCT 2026"},
            {"surface":"Site tags, pixels and event firing","status":"NOT MEASURED","method":"No rendered-browser, tag-manager, network or event-firing test was performed; no tag or pixel is described as absent.","date":"06 OCT 2026"},
            {"surface":"Spend, ROAS, CPA and conversion outcomes","status":"UNAVAILABLE — private to advertiser","method":"No public account, analytics, booking or order data; no estimate, benchmark or projection substituted.","date":"06 OCT 2026"}
        ]
        notes = lead.pop("email_notes")
        subjects = lead.pop("email_subjects")
        signature = "— Smart Pursuit\nsmartpursuit3@gmail.com | 7095024220"
        route = f"The published route, {lead['contact']['email']}, is used only to ask for the right owner; I have not assumed it is the paid-media team."
        email_bodies = [
            f"Hi {lead['brand']} team,\n\n{notes[0]}\n\nThe public record does not expose private spend, orders or event firing. {route}\n\nWould you point me to the person responsible for paid-media measurement?\n\n{signature}",
            f"Hi {lead['brand']} team,\n\n{notes[1]}\n\nI have not converted a public ad count, domain query or click path into a performance claim.\n\nWould a short source-linked handoff be useful?\n\n{signature}",
            f"Hi {lead['brand']} team,\n\n{notes[2]}\n\nThat is an account-side verification step only, not a claim that anything is missing or underperforming.\n\nIs this the right team to ask?\n\n{signature}",
            f"Hi {lead['brand']} team,\n\n{notes[3]}\n\nThis is my last note. If the published route is not the right one, no reply is needed. No private result has been inferred.\n\n{signature}"
        ]
        lead["outreach"] = [[subject, body] for subject, body in zip(subjects, email_bodies)]
        days = ["Day 1", "Day 3", "Day 7", "Day 14"]
        lead["outreach_sequence"] = [
            {"day": day, "subject": subject, "body": body, "source_ids": list(source_ids)}
            for day, (subject, body), source_ids in zip(days, lead["outreach"], lead["email_source_ids"])
        ]

    for lead in leads:
        assert len(lead["outreach"]) == 4
        assert [item["day"] for item in lead["outreach_sequence"]] == ["Day 1", "Day 3", "Day 7", "Day 14"]
        assert sum(int(x["score"]) for x in lead["score_breakdown"]) == int(lead["score_20"]), f"Score does not sum: {lead['brand']}"
        assert float(lead["score"]) == float(lead["score_20"]) / 2, f"Score scale mismatch: {lead['brand']}"
        source_ids = {s["id"] for s in lead["sources_detail"]}
        for ad in lead["ad_evidence"]:
            assert ad.get("library_id") and ad.get("page_id") and ad.get("source_id") in source_ids, f"Missing direct Meta evidence: {lead['brand']}"
            assert ad.get("status", "").startswith("Active"), f"Selected Meta record not active: {lead['brand']}"
        for subject, body in lead["outreach"]:
            assert "smartpursuit3@gmail.com" in body and "7095024220" in body, lead["brand"]
    return data

def check_deduplication(leads):
    registry = json.loads(REGISTRY.read_text(encoding="utf-8"))
    known_names = {norm_name(k) for k in registry.get("brands", {})}
    known_domains = {norm_domain(v.get("website", "")) for v in registry.get("brands", {}).values()}
    existing_names, existing_domains = set(), set()
    for folder in (REPO / "niches").glob("*"):
        if not folder.is_dir():
            continue
        for path in (folder / "leads").glob("*.json"):
            try:
                obj = json.loads(path.read_text(encoding="utf-8"))
            except Exception as exc:
                raise ValueError(f"Unreadable existing lead file {path}: {exc}")
            # No existing lead file is exempt, even if it appears to match a prior
            # interrupted Batch 7 output. This one-time builder must fail closed.
            if obj.get("brand"):

                existing_names.add(norm_name(obj["brand"]))
            if obj.get("website"):
                existing_domains.add(norm_domain(obj["website"]))
    # CSV inventories are also reviewed, not just the central registry and JSON packs.
    # This includes niche lead indexes, one-sheet exports, and archived discovery CSVs.
    name_aliases = {"brand", "brandname", "company", "companyname", "business", "businessname", "lead", "leadname", "name", "advertiser"}
    domain_aliases = {"website", "domain", "url", "companyurl", "businessurl", "homepage", "rootdomain"}
    for path in (REPO / "niches").rglob("*.csv"):
        try:
            with path.open("r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    continue
                names = [h for h in reader.fieldnames if norm_name(h) in name_aliases]
                domains = [h for h in reader.fieldnames if norm_name(h) in domain_aliases]
                for row in reader:
                    for col in names:
                        value = str(row.get(col, "") or "").strip()
                        if value:
                            existing_names.add(norm_name(value))
                            break
                    for col in domains:
                        value = str(row.get(col, "") or "").strip()
                        if value:
                            existing_domains.add(norm_domain(value))
                            break
        except Exception as exc:
            raise ValueError(f"Unreadable existing niche CSV {path}: {exc}")
    # Niche-local registries may contain additional brands not represented in a
    # JSON lead file or current CSV export, so include them independently too.
    for path in (REPO / "niches").glob("*/data/seen_leads.json"):
        try:
            local_registry = json.loads(path.read_text(encoding="utf-8"))
            for name, record in local_registry.get("brands", {}).items():
                existing_names.add(norm_name(name))
                if isinstance(record, dict) and record.get("website"):
                    existing_domains.add(norm_domain(record["website"]))
        except Exception as exc:
            raise ValueError(f"Unreadable niche registry {path}: {exc}")
    duplicates = []
    for lead in leads:
        name_key, domain_key = norm_name(lead["brand"]), norm_domain(lead["website"])
        if name_key in known_names or domain_key in known_domains or name_key in existing_names or domain_key in existing_domains:
            duplicates.append((lead["brand"], lead["website"]))
    if duplicates:
        raise ValueError(f"Deduplication failed: {duplicates}")
    return registry


def make_outreach_markdown(lead):
    c = lead["contact"]
    lines = [
        f"# {lead['brand']} — Outreach templates",
        "",
        f"> Batch 7 · drafted from public sources checked 06 Oct 2026. Internal outreach copy; verify the source links again before sending.",
        "> The published route below is a routing contact, not a confirmed media buyer. The sequence makes no private performance claim.",
        "",
        f"**To (published first-party route):** `{c['email']}`<br>",
        f"**Contact role:** {c.get('route_role', 'Published general contact; buyer not confirmed')}<br>",
        f"**Provenance:** {lead['contact_source']}",
        "",
        "---",
        "",
    ]
    days = ["Day 1", "Day 3", "Day 7", "Day 14"]
    srcmap = {source["id"]: source for source in lead.get("sources_detail", [])}
    for i, (subject, body) in enumerate(lead["outreach"]):
        ids = lead.get("email_source_ids", [[]] * 4)[i]
        evidence_links = " · ".join(f"[{sid}]({srcmap[sid]['url']})" for sid in ids if sid in srcmap)
        lines.extend([
            f"## Email {i + 1} ({days[i]}): {subject}",
            "",
            "```text",
            body.strip(),
            "```",
            "",
            f"**Evidence to recheck:** {evidence_links}",
            "",
        ])
    lines.extend([
        "---",
        "",
        "**Before sending:** re-open the linked ad/source, confirm the offer or creative is still current, and ask for the right owner rather than assuming this route is a buyer.",
        "",
    ])
    return "\n".join(lines)


def write_json_and_outreach(leads):
    for lead in leads:
        (ROOT / "leads" / f"{lead['slug']}.json").write_text(
            json.dumps(lead, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        (ROOT / "outreach" / f"{lead['slug']}-outreach-templates.md").write_text(
            make_outreach_markdown(lead), encoding="utf-8"
        )


def _story_value(s):
    return " ".join(str(s).replace("\n", " ").split())


def _google_card_text(lead):
    status = lead["google_evidence"]["status"].casefold()
    brand = lead.get("brand", "").casefold()
    if "no domain-query results" in status or "zero domain-query results" in status:
        return "0 domain results", "Not proof of no activity"
    if brand == "go desi":
        return "Brand tie unconfirmed", "Opened Haystack agency creative; GO DESi not named"
    if brand == "deliciae":
        return "Archive creative noted", "Last shown 05 Oct; removed, not current status"
    if "la pino" in brand:
        return "Archive creative noted", "Last shown 05 Oct; not labelled Active"
    return "Checked; status limited", "See individual record and last-shown caveat"

def _local_card_text(lead):
    m = lead.get("market_signal", "")
    local_status = lead.get("local_evidence_status", "")
    if "partially verified" in local_status.casefold():
        return "Indexed; direct fetch failed", "City detail needs live recheck"
    if "unconfirmed" in m.casefold() and any(x in m.casefold() for x in ("bengaluru", "bangalore", "blr")):
        return "Outlet unconfirmed", "Selected creative names Bengaluru"
    ad_copy = " ".join(a.get("creative_excerpt", "") for a in lead.get("ad_evidence", [])).casefold()
    if any(x in ad_copy for x in ("bengaluru", "bangalore", "blr")):
        return "City in selected creative", "No audience targeting inferred"
    if any(k in m.casefold() for k in ("bengaluru address", "bengaluru store", "bengaluru corporate office", "bengaluru location", "bangalore address", "bangalore delivery")):
        return "First-party local context", "No ad-targeting inference"
    return "Local context published", "Source and limits noted in report"


def _register_fallback_fonts():
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        sans = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        mono = Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")
        if sans.exists() and "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(sans)))
        if mono.exists() and "DejaVuSansMono" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVuSansMono", str(mono)))
    except Exception:
        pass


def pdf_safe(text, mono=False):
    text = html.escape(str(text), quote=False).replace("\n", "<br/>")
    replacement = "DejaVuSansMono" if mono else "DejaVuSans"
    text = text.replace("₹", f"<font name='{replacement}'>₹</font>")
    return text


def make_pdf(lead, total_pages=5, destination=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import (
        BaseDocTemplate, PageTemplate, Frame, Flowable, Paragraph, Spacer, Table,
        TableStyle, KeepTogether, CondPageBreak, NextPageTemplate, PageBreak,
        HRFlowable,
    )
    from reportlab.lib.utils import ImageReader

    _register_fallback_fonts()
    page_w, page_h = letter
    navy = colors.HexColor("#" + NAVY)
    navy2 = colors.HexColor("#" + NAVY_2)
    teal = colors.HexColor("#" + TEAL)
    gold = colors.HexColor("#" + GOLD)
    panel = colors.HexColor("#" + PANEL)
    rule = colors.HexColor("#" + RULE)
    muted = colors.HexColor("#" + MUTED)
    ink = colors.HexColor("#" + INK)
    white = colors.white
    logo = ROOT / "assets" / "smart-pursuit-logo.png"

    class CoverMarker(Flowable):
        def __init__(self):
            super().__init__()
            self.width = 1
            self.height = 1
        def draw(self):
            pass

    class AuditDoc(BaseDocTemplate):
        def __init__(self, path, total):
            super().__init__(path, pagesize=letter, leftMargin=42, rightMargin=42,
                             topMargin=84, bottomMargin=52, title=f"{lead['brand']} — Paid Media & Measurement Audit",
                             author="Smart Pursuit")
            self.total = total
            cover_frame = Frame(42, 50, page_w - 84, page_h - 100, id="cover",
                                leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
            body_frame = Frame(42, 58, page_w - 84, 650, id="body",
                               leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
            self.addPageTemplates([
                PageTemplate(id="Cover", frames=[cover_frame], onPage=self._cover),
                PageTemplate(id="Body", frames=[body_frame], onPage=self._body),
            ])

        def _cover(self, canv, doc):
            canv.saveState()
            canv.setFillColor(navy)
            canv.rect(0, 0, page_w, page_h, fill=1, stroke=0)
            canv.setFillColor(gold)
            canv.rect(36, 628, 4, 118, fill=1, stroke=0)
            if logo.exists():
                canv.drawImage(str(logo), 55, 666, width=78, height=78, preserveAspectRatio=True, mask="auto")
            canv.setFillColor(gold)
            canv.setFont("Helvetica-Bold", 8)
            canv.drawString(150, 728, "PAID MEDIA & MEASUREMENT AUDIT")
            canv.setFillColor(white)
            brand_size = 27
            while stringWidth(lead["brand"], "Helvetica-Bold", brand_size) > 404 and brand_size > 19:
                brand_size -= 1
            canv.setFont("Helvetica-Bold", brand_size)
            canv.drawString(150, 690, lead["brand"])
            subtitle = Paragraph(pdf_safe(lead.get("headline", "Public-evidence snapshot")), ParagraphStyle(
                "CoverSubtitle", fontName="Helvetica", fontSize=9.2, leading=12.2,
                textColor=colors.HexColor("#D7E2EC"), spaceAfter=0,
            ))
            sw, sh = subtitle.wrap(408, 46)
            subtitle.drawOn(canv, 150, max(643, 654 - sh))
            canv.setFillColor(colors.HexColor("#C5D1DC"))
            canv.setFont("Helvetica-Bold", 7.2)
            canv.drawString(55, 622, "BATCH 7  ·  PUBLIC-EVIDENCE SNAPSHOT  ·  VERIFIED 06 OCT 2026")

            google_value, google_note = _google_card_text(lead)
            local_value, local_note = _local_card_text(lead)
            cards = [
                ("SELECTED META RECORDS", f"{len(lead['ad_evidence'])} selected", "Active on the capture date; not an account-wide total"),
                ("GOOGLE ADS TRANSPARENCY", google_value, google_note),
                ("BENGALURU RELEVANCE", local_value, local_note),
                ("PUBLIC-EVIDENCE FIT", f"{lead['score_20']} / 20", f"{lead['score']} / 10 · triage only, not performance"),
            ]
            card_w, card_h = 252, 74
            positions = [(42, 525), (318, 525), (42, 435), (318, 435)]
            for (label, value, note), (x, y) in zip(cards, positions):
                canv.setFillColor(colors.HexColor("#F4F7FA"))
                canv.roundRect(x, y, card_w, card_h, 7, fill=1, stroke=0)
                canv.setFillColor(gold)
                canv.roundRect(x, y, 4, card_h, 2, fill=1, stroke=0)
                canv.setFillColor(muted)
                canv.setFont("Helvetica-Bold", 6.7)
                canv.drawString(x + 14, y + 56, label)
                value_size = 14 if len(value) < 23 else 11
                canv.setFillColor(navy)
                canv.setFont("Helvetica-Bold", value_size)
                canv.drawString(x + 14, y + 35, value[:42])
                note_para = Paragraph(pdf_safe(note), ParagraphStyle(
                    "CardNote", fontName="Helvetica", fontSize=7.0, leading=8.6,
                    textColor=muted,
                ))
                nw, nh = note_para.wrap(card_w - 27, 20)
                note_para.drawOn(canv, x + 14, y + 7)

            # Boundary statement is intentionally explicit; the cover is not a performance claim.
            canv.setFillColor(colors.HexColor("#18334B"))
            canv.roundRect(42, 350, 528, 56, 7, fill=1, stroke=0)
            canv.setFillColor(gold)
            canv.setFont("Helvetica-Bold", 7.2)
            canv.drawString(56, 385, "SCOPE BOUNDARY")
            boundary = Paragraph(
                "Public pages and selected ad-library records only. Private account spend, ROAS, CPA, conversions, bookings and event firing were not accessed; unmeasured values are not estimated.",
                ParagraphStyle("Boundary", fontName="Helvetica", fontSize=8.0, leading=10.3,
                               textColor=colors.HexColor("#E1EAF1")),
            )
            bw, bh = boundary.wrap(496, 30)
            boundary.drawOn(canv, 56, 360)

            # Prepared-for and agency contact boxes.
            canv.setFillColor(colors.HexColor("#F4F7FA"))
            canv.roundRect(42, 213, 252, 104, 7, fill=1, stroke=0)
            canv.setFillColor(gold)
            canv.roundRect(42, 213, 4, 104, 2, fill=1, stroke=0)
            canv.setFillColor(muted)
            canv.setFont("Helvetica-Bold", 7.2)
            canv.drawString(57, 294, "PREPARED FOR")
            canv.setFillColor(navy)
            canv.setFont("Helvetica-Bold", 15)
            canv.drawString(57, 270, (lead["brand"] + " team")[:36])
            canv.setFont("Helvetica", 8)
            canv.setFillColor(ink)
            canv.drawString(57, 250, "Bengaluru food-lead batch · 07")
            canv.drawString(57, 234, "Source-checked 06 October 2026")

            canv.setFillColor(navy2)
            canv.roundRect(318, 213, 252, 104, 7, fill=1, stroke=0)
            canv.setStrokeColor(gold)
            canv.setLineWidth(1)
            canv.roundRect(318, 213, 252, 104, 7, fill=0, stroke=1)
            canv.setFillColor(gold)
            canv.setFont("Helvetica-Bold", 7.2)
            canv.drawString(333, 294, "CONTACT · SMART PURSUIT")
            canv.setFillColor(white)
            canv.setFont("Helvetica-Bold", 9.2)
            canv.drawString(333, 270, "smartpursuit3@gmail.com")
            canv.setFont("Helvetica", 8.5)
            canv.drawString(333, 250, "7095024220")
            canv.setFillColor(colors.HexColor("#D7E2EC"))
            canv.setFont("Helvetica", 7)
            canv.drawString(333, 231, "Brand-published route is listed inside the report.")

            canv.setStrokeColor(colors.HexColor("#294257"))
            canv.line(42, 54, 570, 54)
            canv.setFillColor(colors.HexColor("#B7C6D3"))
            canv.setFont("Helvetica", 7.2)
            canv.drawString(42, 35, "Smart Pursuit  ·  smartpursuit3@gmail.com  ·  7095024220")
            canv.drawRightString(570, 35, f"Page 1 of {self.total}")
            canv.restoreState()

        def _body(self, canv, doc):
            canv.saveState()
            canv.setFillColor(white)
            canv.rect(0, 0, page_w, page_h, fill=1, stroke=0)
            if logo.exists():
                canv.drawImage(str(logo), 40, 744, width=40, height=40, preserveAspectRatio=True, mask="auto")
            canv.setFillColor(navy)
            canv.setFont("Helvetica-Bold", 7.4)
            canv.drawString(91, 766, "FOOD / BENGALURU  ·  PUBLIC-EVIDENCE AUDIT")
            canv.setFillColor(muted)
            canv.setFont("Helvetica", 7)
            canv.drawRightString(570, 766, f"{lead['brand']}  ·  06 OCT 2026")
            canv.setStrokeColor(rule)
            canv.setLineWidth(0.8)
            canv.line(42, 735, 570, 735)
            canv.line(42, 39, 570, 39)
            canv.setFillColor(muted)
            canv.setFont("Helvetica", 6.8)
            canv.drawString(42, 24, "Smart Pursuit  ·  smartpursuit3@gmail.com  ·  7095024220")
            canv.drawRightString(570, 24, f"Page {doc.page} of {self.total}")
            canv.restoreState()

    # Styles are intentionally compact but readable to target five pages.
    S = {
        "section": ParagraphStyle("Section", fontName="Helvetica-Bold", fontSize=11.0, leading=14,
                                   textColor=navy, spaceBefore=5, spaceAfter=5),
        "body": ParagraphStyle("Body", fontName="Helvetica", fontSize=8.0, leading=10.3,
                                textColor=ink, spaceAfter=3),
        "small": ParagraphStyle("Small", fontName="Helvetica", fontSize=7.25, leading=9.1,
                                 textColor=ink),
        "tiny": ParagraphStyle("Tiny", fontName="Helvetica", fontSize=6.7, leading=8.2,
                                textColor=ink),
        "thead": ParagraphStyle("TableHeader", fontName="Helvetica-Bold", fontSize=6.7, leading=8.2,
                                 textColor=white),
        "label": ParagraphStyle("Label", fontName="Helvetica-Bold", fontSize=7.0, leading=8.5,
                                 textColor=navy),
        "muted": ParagraphStyle("Muted", fontName="Helvetica", fontSize=7.1, leading=8.8,
                                 textColor=muted),
        "findingtitle": ParagraphStyle("FindingTitle", fontName="Helvetica-Bold", fontSize=9.0,
                                        leading=11.2, textColor=navy),
        "evidence": ParagraphStyle("Evidence", fontName="Courier", fontSize=7.0, leading=8.8,
                                    textColor=ink, splitLongWords=1),
        "italic": ParagraphStyle("Italic", fontName="Helvetica-Oblique", fontSize=7.65, leading=9.8,
                                 textColor=muted),
        "callout": ParagraphStyle("Callout", fontName="Helvetica", fontSize=7.8, leading=10.0,
                                   textColor=ink),
        "contents": ParagraphStyle("Contents", fontName="Helvetica", fontSize=7.6, leading=9.2,
                                    textColor=ink),
        "nexttitle": ParagraphStyle("NextTitle", fontName="Helvetica-Bold", fontSize=8.6,
                                     leading=10.2, textColor=navy),
        "next": ParagraphStyle("Next", fontName="Helvetica", fontSize=7.7, leading=9.8,
                               textColor=ink),
        "score": ParagraphStyle("Score", fontName="DejaVuSans", fontSize=7.8, leading=9.5,
                                 textColor=navy),
    }

    def P(text, style="body"):
        return Paragraph(pdf_safe(text), S[style])

    def heading(number, title):
        return Paragraph(
            f"<font color='#{TEAL}'>{html.escape(number)}</font>  <b>{html.escape(title.upper())}</b>",
            S["section"],
        )

    def source_link(url, label="Open source"):
        safe_url = html.escape(str(url), quote=True).replace("&amp;", "&amp;")
        return Paragraph(f"<link href=\"{safe_url}\"><font color=\"#{TEAL}\"><u>{html.escape(label)}</u></font></link>", S["tiny"])

    def status_color(status):
        s = status.upper()
        if "UNAVAILABLE" in s:
            return colors.HexColor("#E2E8F0")
        if "NOT MEASURED" in s or "UNCONFIRMED" in s:
            return colors.HexColor("#FEF3C7")
        return colors.HexColor("#D1FAE5")

    def section_table(rows, widths, header=True, font_style="tiny", repeatRows=1):
        table = Table(rows, colWidths=widths, repeatRows=(repeatRows if header else 0), hAlign="LEFT")
        commands = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ]
        if header:
            commands += [
                ("BACKGROUND", (0, 0), (-1, 0), navy),
                ("TEXTCOLOR", (0, 0), (-1, 0), white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        table.setStyle(TableStyle(commands))
        return table

    story = [CoverMarker(), NextPageTemplate("Body"), PageBreak()]

    # CONTENTS + METHOD
    story.append(heading("", "Contents"))
    contents = [
        [P("01  Method", "contents"), P("02  The ad account + coverage", "contents")],
        [P("03  Findings", "contents"), P("03b  Evidence register", "contents")],
        [P("04  Opportunity score", "contents"), P("05  Next steps", "contents")],
    ]
    toc = Table(contents, colWidths=[264, 264], hAlign="LEFT")
    toc.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), panel),
        ("BOX", (0, 0), (-1, -1), 0.45, rule),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.extend([toc, Spacer(1, 6), heading("01", "Method"), P(
        "Dated public-evidence snapshot checked 06 October 2026. Meta evidence is cited by individual Library ID and advertiser/page ID; the selected cards are not an account-wide count. Public ad counts and rendered variants can change daily; recheck before relying on them. Meta search-type parameters can be ignored by the interface, so keyword results are not treated as page or advertiser proof. Bengaluru relevance is based only on explicit creative or first-party material; no audience targeting is inferred.", "body"), P(
        "Contacts and offer details are taken from the brand's own website, help or corporate pages and remain routing contacts unless a buyer role is explicitly published. No private account, booking, order or analytics records were accessed. Spend, ROAS, CPA, conversion outcomes and event firing are marked unavailable or not measured. A public ad card or URL alone does not prove that tracking events fire.", "body"),
    ])

    # THE AD ACCOUNT: concise label/value table.
    story.append(heading("02", "The ad account"))
    meta_summary = "; ".join(
        f"ID {a['library_id']} (started {a.get('started_on', 'not published')})"
        for a in lead.get("ad_evidence", [])
    )
    contact = lead["contact"]
    account_rows = [
        [P("Brand / website", "label"), P(f"{lead['brand']} · {lead['website']}", "small")],
        [P("Local context", "label"), P(lead.get("market_signal", "Not established from reviewed public sources"), "small")],
        [P("Selected Meta records", "label"), P(f"{meta_summary}. Active on 06 Oct 2026; selected records only.", "small")],
        [P("Google Ads", "label"), P(f"{lead['google_evidence']['status']}: {lead['google_evidence']['detail']}", "small")],
        [P("Published contact", "label"), P(f"{contact.get('email','')} · {contact.get('route_role','')} {lead.get('contact_source','')}", "small")],
        [P("Public-evidence fit", "label"), P(f"{lead['score_20']} / 20 · {lead['score']} / 10 · {lead['score_band']}. {lead.get('score_note','')}", "small")],
    ]
    at = Table(account_rows, colWidths=[118, 410], hAlign="LEFT")
    at.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), panel),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(at)
    story.append(Spacer(1, 5))

    story.append(P("COVERAGE · each line is dated 06 OCT 2026", "label"))
    cov_rows = [[P("Surface", "thead"), P("Status", "thead"), P("Method / boundary", "thead")]]
    for c in lead.get("coverage", []):
        cov_rows.append([
            P(c["surface"], "tiny"), P(c["status"], "tiny"), P(c["method"], "tiny")
        ])
    cov = Table(cov_rows, colWidths=[112, 132, 284], repeatRows=1, hAlign="LEFT")
    cov_commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for r, c in enumerate(lead.get("coverage", []), start=1):
        cov_commands.append(("BACKGROUND", (1, r), (1, r), status_color(c["status"])))
    cov.setStyle(TableStyle(cov_commands))
    boundary_rows = [
        [P("AD LIBRARY", "label"), P("Selected Library IDs and page identities establish the visible record and copy only; no account-wide inventory or audience targeting is inferred.", "tiny")],
        [P("GOOGLE", "label"), P(lead["google_evidence"]["detail"], "tiny")],
        [P("LOCAL", "label"), P("City relevance comes from the selected creative or first-party material; it does not establish geographic targeting.", "tiny")],
        [P("PRIVATE DATA", "label"), P("Spend, ROAS, CPA, conversion outcomes and event firing are unavailable/not measured unless a source explicitly establishes them.", "tiny")],
    ]
    boundary_table = Table(boundary_rows, colWidths=[98, 430], hAlign="LEFT")
    boundary_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8F4F4")),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.extend([cov, Spacer(1, 5), P("EVIDENCE BOUNDARIES", "label"), boundary_table, Spacer(1, 7)])

    # FINDINGS: title+raw evidence atomic, analysis and next check can flow.
    story.append(CondPageBreak(250))
    story.append(heading("03", "Findings"))
    sev_colors = {
        "CRITICAL": (colors.HexColor("#FEE2E2"), colors.HexColor("#DC2626")),
        "HIGH": (colors.HexColor("#FFEDD5"), colors.HexColor("#EA580C")),
        "MEDIUM": (colors.HexColor("#FEF3C7"), colors.HexColor("#B45309")),
        "LOW": (colors.HexColor("#D1FAE5"), colors.HexColor("#047857")),
        "NOTE": (colors.HexColor("#E2E8F0"), colors.HexColor("#475569")),
    }
    for idx, f in enumerate(lead.get("findings", []), start=1):
        severity = str(f.get("priority", "NOTE")).upper()
        bg, fg = sev_colors.get(severity, sev_colors["NOTE"])
        badge = Table(
            [[P(severity, "label"), P(f"{idx:02d}  {f['title']}", "findingtitle")]],
            colWidths=[68, 460], hAlign="LEFT",
        )
        badge.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), bg),
            ("BACKGROUND", (1, 0), (1, 0), panel),
            ("TEXTCOLOR", (0, 0), (0, 0), fg),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOX", (0, 0), (-1, -1), 0.45, rule),
            ("INNERGRID", (0, 0), (-1, -1), 0.35, rule),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        evidence_text = (
            f"source_ids: {', '.join(f.get('source_ids', []))}\n"
            f"observed: {f.get('evidence', '')}\n"
            f"checked: 06 OCT 2026"
        )
        evidence_box = Table([[Paragraph(pdf_safe(evidence_text, mono=True), S["evidence"])]], colWidths=[528])
        evidence_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("BOX", (0, 0), (-1, -1), 0.4, rule),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(KeepTogether([badge, evidence_box]))
        story.append(Spacer(1, 3))
        story.append(P(f["analysis"], "italic"))
        story.append(P(f"NEXT CHECK · {f['next_check']}", "small"))
        story.append(Spacer(1, 6))

    # EVIDENCE REGISTER: all cited source records, including first-party contact provenance.
    story.extend([CondPageBreak(100), heading("03b", "Evidence register")])
    ev_rows = [[P("ID / source", "thead"), P("Value observed", "thead"), P("Date", "thead"), P("Link", "thead")]]
    for src in lead.get("sources_detail", []):
        ev_rows.append([
            P(f"{src['id']} · {src['label']}", "tiny"),
            P(src.get("observed", ""), "tiny"),
            P(src.get("verified_on", "2026-10-06"), "tiny"),
            source_link(src["url"], f"Open {src['id']}"),
        ])
    evt = Table(ev_rows, colWidths=[108, 300, 62, 58], repeatRows=1, hAlign="LEFT")
    evt.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [white, panel]),
    ]))
    story.extend([evt, Spacer(1, 8)])

    # SCORE: public evidence fit only; this is not an efficacy or ROI rating.
    story.extend([CondPageBreak(240), heading("04", "Opportunity score")])
    score_rows = [[P("Dimension", "thead"), P("Score", "thead"), P("Basis", "thead")]]
    for item in lead.get("score_breakdown", []):
        score_rows.append([
            P(item["dimension"], "tiny"),
            P(f"{'●' * int(item['score'])}{'○' * (4 - int(item['score']))}  {item['score']} / 4", "score"),
            P(item["basis"], "tiny"),
        ])
    sc = Table(score_rows, colWidths=[148, 52, 328], repeatRows=1, hAlign="LEFT")
    sc.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [white, panel]),
    ]))
    story.append(sc)
    total_score = Table([[P(f"OVERALL PUBLIC-EVIDENCE FIT  ·  {lead['score_20']} / 20  ·  {lead['score']} / 10  ·  {lead['score_band']}", "score")]], colWidths=[528])
    total_score.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E8F4F4")),
        ("BOX", (0, 0), (-1, -1), 0.8, teal),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    qualification = Table([[P(
        "QUALIFICATION · Read before acting. This is a triage score of public paid signal, local relevance, public CTA visibility, first-party source coverage and contact-route relevance. It is not an ad-performance, revenue, sales-probability, spend, forecast or ROI score. Private results were not accessed. Scores are not directly comparable with the legacy batches' prior scoring framework.", "callout")]], colWidths=[528])
    qualification.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF8E7")),
        ("BOX", (0, 0), (-1, -1), 0.5, gold),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    account_check_rows = [
        [P("Outcome definition", "label"), P("Name the event or business outcome, its counting rule and any deduplication rule before comparing systems.", "tiny")],
        [P("Time window", "label"), P("Use the same date range, timezone and attribution window across any records being reconciled.", "tiny")],
        [P("Path mapping", "label"), P("Trace the selected Library ID through the visible click route to the brand's own booking, enquiry or order record where applicable.", "tiny")],
        [P("Reconciliation", "label"), P("Keep platform, analytics and CRM/order data side by side until event definitions and windows are comparable; do not add incompatible totals.", "tiny")],
    ]
    account_check = Table(account_check_rows, colWidths=[112, 416], hAlign="LEFT")
    account_check.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), panel),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    score_key_rows = [
        [P("Direct paid signal", "label"), P("A page-resolved active public ad record; not spend or outcome.", "tiny")],
        [P("Bengaluru relevance", "label"), P("Explicit local creative or first-party context; not audience targeting.", "tiny")],
        [P("Public CTA visibility", "label"), P("A visible public next step or destination; not a conversion.", "tiny")],
        [P("First-party coverage", "label"), P("Breadth of official product, store, contact or corporate evidence reviewed.", "tiny")],
        [P("Contact-route relevance", "label"), P("How suitable the published route is for a referral request; no buyer identity is inferred.", "tiny")],
    ]
    score_key = Table(score_key_rows, colWidths=[155, 373], hAlign="LEFT")
    score_key.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (0, -1), panel),
        ("GRID", (0, 0), (-1, -1), 0.35, rule),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.extend([Spacer(1, 5), P("SCORING KEY", "label"), score_key, Spacer(1, 6),
                  P("ACCOUNT-SIDE VALIDATION CHECKLIST", "label"), account_check, Spacer(1, 8)])

    # NEXT STEPS + CTA + close; suggested effort labels are operational, not outcome estimates.
    story.extend([CondPageBreak(200), heading("05", "Next steps"), P(
        "Suggested effort labels are planning aids only. First re-open the dated public sources; then use internal account/CRM/order records only if the company chooses to share them.", "muted")])
    ad_ids = ", ".join(a["library_id"] for a in lead.get("ad_evidence", []))
    meta_dates = ", ".join(a.get("started_on", "date not exposed") for a in lead.get("ad_evidence", []))
    step_specs = [
        ("LOW EFFORT", "Route the question", f"Use {contact.get('email','the published contact')} as a routing contact only; ask for the paid-media measurement owner. The published source does not identify this inbox as a buyer."),
        ("UNDER AN HOUR", "Re-open the selected card", f"Recheck Library ID(s) {ad_ids}; the selected records showed start date(s) {meta_dates}. Confirm today's status, page identity, copy, CTA and any currently exposed destination before reusing the observation."),
        ("UNDER AN HOUR", "Complete the finding-specific public check", " ".join(f["next_check"] for f in lead.get("findings", [])[:2])),
        ("HALF A DAY", "Reconcile only with authorized private records", lead.get("findings", [{}])[-1].get("next_check", "") + " If access is authorized, align the event definition, date window, timezone, Library ID and the corresponding booking/order or CRM record. Do not infer an outcome from the public ad card."),
    ]
    for step_index, (effort, title, text) in enumerate(step_specs):
        # A conditional break preserves complete, actionable step cards; the final card and
        # source handoff can continue on page 5 when the prior page is full.
        if step_index == 3:
            story.append(CondPageBreak(240))
        # This is intentional ReportLab markup; P() escapes markup for ordinary text.
        badge = Paragraph(f"<b>{html.escape(effort)}</b>", S["tiny"])
        content = [P(title, "nexttitle"), P(text, "next")]
        row = Table([[badge, content]], colWidths=[100, 428], hAlign="LEFT")
        row.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#FEF3C7")),
            ("BACKGROUND", (1, 0), (1, 0), panel),
            ("BOX", (0, 0), (-1, -1), 0.4, rule),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([row, Spacer(1, 5)])

    # Compact, source-linked handoff: makes the final page useful as a practical re-check sheet.
    contact_src = next((s for s in lead.get("sources_detail", []) if s["id"] in lead["contact_source_detail"]["source_ids"]), None)
    handoff_route = Paragraph(
        f"{pdf_safe(contact.get('email', ''))} · {pdf_safe(contact.get('phone', ''))}<br/>{pdf_safe(contact.get('route_role', ''))}",
        S["small"],
    )
    handoff_contact = Table([[P("PUBLISHED ROUTE", "label"), handoff_route, source_link(contact_src["url"], "Verify route") if contact_src else P("Source listed in register", "tiny")]], colWidths=[100, 350, 78], hAlign="LEFT")
    handoff_contact.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#E8F4F4")),
        ("BACKGROUND", (1, 0), (-1, -1), panel),
        ("GRID", (0, 0), (-1, -1), 0.4, rule),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([Spacer(1, 4), P("SOURCE LINKS TO REOPEN BEFORE OUTREACH", "label")])
    link_rows = [[P("ID", "thead"), P("Source", "thead"), P("Open", "thead")]]
    for src in lead.get("sources_detail", []):
        source_summary = f"{pdf_safe(src['label'])}<br/><font color='#{MUTED}'>{pdf_safe(src.get('observed',''))}</font>"
        link_rows.append([P(src["id"], "tiny"), Paragraph(source_summary, S["tiny"]), source_link(src["url"], "Open source")])
    link_table = Table(link_rows, colWidths=[42, 404, 82], repeatRows=1, hAlign="LEFT")
    link_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("GRID", (0, 0), (-1, -1), 0.35, rule), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    handoff_rows = [
        [P("EVENT", "label"), P("Name the agreed outcome and its counting/deduplication rule.", "tiny")],
        [P("IDENTITY", "label"), P("Map the selected Library ID or campaign key to internal records.", "tiny")],
        [P("WINDOW", "label"), P("Align dates, timezone and attribution window.", "tiny")],
        [P("SYSTEM", "label"), P("Name the source-of-record system and authorized owner.", "tiny")],
    ]
    handoff_check = Table(handoff_rows, colWidths=[72, 456], hAlign="LEFT")
    handoff_check.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8F4F4")),
        ("GRID", (0, 0), (-1, -1), 0.35, rule), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    send_check_rows = [
        [P("DATE", "label"), P("Re-open each cited Library ID and update the status/date if the public card has changed.", "tiny")],
        [P("CLAIM", "label"), P("Use only the cited copy or first-party terms; do not infer audience geography, tracking events or private outcomes.", "tiny")],
        [P("ROUTE", "label"), P("Use the published first-party email as a routing contact and ask for the right owner if needed.", "tiny")],
    ]
    send_check = Table(send_check_rows, colWidths=[62, 466], hAlign="LEFT")
    send_check.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#FEF3C7")),
        ("GRID", (0, 0), (-1, -1), 0.35, rule), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.extend([handoff_contact, Spacer(1, 6), link_table, Spacer(1, 6),
                  P("IF AN ACCOUNT-SIDE REVIEW IS AUTHORIZED", "label"), handoff_check, Spacer(1, 6),
                  P("OUTREACH CHECK BEFORE SENDING", "label"), send_check, Spacer(1, 6)])

    cta = Table([[Paragraph("<font color='#FFFFFF'><b>Happy to walk through any of this on a short call.</b></font>", S["findingtitle"]),
                  Paragraph("<font color='#FFFFFF'><b>GET IN TOUCH</b><br/>smartpursuit3@gmail.com<br/>7095024220</font>", S["small"])]], colWidths=[320, 208])
    cta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), navy),
        ("TEXTCOLOR", (0, 0), (-1, -1), white),
        ("BOX", (0, 0), (-1, -1), 0.5, gold),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    close = Table([[P(
        "ABOUT THIS REPORT · Smart Pursuit prepared this dated snapshot from the public sources listed above. The linked sources are the verification route; recheck them because live ad inventory changes. No private account data or client performance result is represented here.", "tiny")]], colWidths=[528])
    close.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E8F4F4")),
        ("BOX", (0, 0), (-1, -1), 0.4, teal),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([Spacer(1, 3), cta, Spacer(1, 5), close])

    doc = AuditDoc(str(destination), total_pages)
    doc.build(story)


def generate_pdfs(leads):
    import pymupdf as fitz
    audit_dir = ROOT / "audits"
    audit_dir.mkdir(exist_ok=True)
    page_counts = {}
    for lead in leads:
        out = audit_dir / f"{lead['slug']}-paid-media-measurement-audit.pdf"
        temp = out.with_suffix(".layout-check.pdf")
        make_pdf(lead, total_pages=99, destination=temp)
        with fitz.open(temp) as pdf:
            page_count = len(pdf)
        if not 5 <= page_count <= 10:
            temp.unlink(missing_ok=True)
            raise RuntimeError(f"{lead['brand']} audit rendered {page_count} pages; expected a readable 5–10 page source-linked report")
        make_pdf(lead, total_pages=page_count, destination=out)
        temp.unlink(missing_ok=True)
        with fitz.open(out) as final_pdf:
            if len(final_pdf) != page_count:
                raise RuntimeError(f"{lead['brand']} final PDF page count changed after layout pass")
            if ROOT.joinpath("assets", "smart-pursuit-logo.png").exists() and not final_pdf[0].get_images(full=True):
                raise RuntimeError(f"{lead['brand']} cover is missing the pack logo")
            for page_number, page in enumerate(final_pdf, start=1):
                page_text = page.get_text()
                if f"Page {page_number} of {page_count}" not in page_text:
                    raise RuntimeError(f"{lead['brand']} PDF footer is missing correct page numbering on page {page_number}")
            if not any(page.get_links() for page in final_pdf):
                raise RuntimeError(f"{lead['brand']} PDF contains no clickable source links")
        page_counts[lead["brand"]] = page_count
        print(f"PDF {out.name}: {page_count} pages, numbered footer, logo and source links verified; {out.stat().st_size:,} bytes")
    return page_counts

def _md_table(text_rows):
    def esc(v):
        return str(v).replace("|", "\\|").replace("\n", " ")
    lines = ["| Claim in the audit | Where to check it | Link | What you should see |",
             "|---|---|---|---|"]
    for row in text_rows:
        lines.append("| " + " | ".join(esc(x) for x in row) + " |")
    return "\n".join(lines)


def extract_step0_b7(existing):
    start = existing.find(STEP0_MARKER)
    if start < 0:
        raise RuntimeError("Batch 7 Step 0 inventory/skill disposition is missing; refusing to rewrite verification history")
    end = existing.find("\n### Batch 7 incremental builder", start)
    if end >= 0:
        return existing[start:end].strip()
    return existing[start:].strip()


def build_verify_batch_section(leads, step0_text):
    data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
    out = [
        "## Batch 7 — Food / Bengaluru (source check: 06 October 2026)\n\n",
        "This is a dated, source-linked public-evidence snapshot. Six leads were retained against a working target of ten; no candidate was added to fill the shortfall. Each selected Meta card was individually opened on 06 Oct 2026 and displayed the named page, Sponsored and Active labels, start date and direct Library ID. One selected card is not an account-wide count and public inventory can change. Google domain counts are any-time discovery snapshots; zero results do not prove no activity. An opened Google creative is attributed only to the named advertiser/visible brand content and is not treated as current-active unless the source says so. Bengaluru relevance comes from the official brand page or selected ad copy, not inferred audience targeting. Spend, ROAS, CPA, orders, bookings, conversion rates, tag/pixel presence and event firing remain unavailable/not measured.\n\n",
    ]
    for lead in leads:
        out.append(f"### {lead['brand']}\n\nSite: `{lead['website']}` · source check 06 October 2026 · Batch 7\n\n")
        out.append(f"**Published contact and provenance:** `{lead['contact']['email']}` — {lead['contact'].get('route_role','')} {lead.get('contact_source_status','')} {lead['contact_source']}\n\n")
        rows = []
        srcmap = {src["id"]: src for src in lead.get("sources_detail", [])}
        for ad in lead.get("ad_evidence", []):
            src = srcmap[ad["source_id"]]
            rows.append((
                f"Meta Ad Library ID {ad['library_id']} — Active and Sponsored; started {ad.get('started_on','not stated')}",
                f"Meta Ad Library · {ad.get('page_name','advertiser page')} · page ID {ad.get('page_id','')}",
                f"[{src['url']}]({src['url']})",
                f"Checked {ad.get('checked_on','')}; direct card and advertiser/page profile: {ad.get('page_profile_url','')}; visible copy: {ad.get('creative_excerpt','')}; CTA: {ad.get('cta','')}; destination boundary: {ad.get('destination_status','')}",
            ))
        for finding in lead.get("findings", []):
            for sid in finding.get("source_ids", []):
                src = srcmap[sid]
                rows.append((finding["title"], f"{src['label']} ({sid})", f"[{src['url']}]({src['url']})",
                             f"Evidence: {finding['evidence']} Analysis boundary: {finding['analysis']} Next check: {finding['next_check']}"))
        for src in lead.get("sources_detail", []):
            rows.append((f"Source register {src['id']} — {src['label']}", src["source_type"],
                         f"[{src['url']}]({src['url']})", src.get("observed", "")))
        g = lead.get("google_evidence", {})
        rows.append((f"Google Ads Transparency status: {g.get('status','not measured')}",
                     "Google Ads Transparency Center · India region · domain and individual creative where available",
                     f"[{g.get('url','')}]({g.get('url','')})", g.get("detail", "")))
        rows.append(("Private performance / tracking measures", "Unavailable to the public reviewer", "—",
                     "Spend, ROAS, CPA, orders/bookings, conversions, tag/pixel presence and event firing were not accessed or estimated. No tag/pixel absence claim is made."))
        rows.append(("Public-evidence fit score", "Five-dimension triage score in report and lead JSON",
                     f"[`leads/{lead['slug']}.json`](leads/{lead['slug']}.json)",
                     f"{lead['score_20']} / 20 ({lead['score']} / 10), triage only; not an ROI, performance, revenue, forecast or sales-probability score."))
        out.append(_md_table(rows))
        out.append("\n\n#### Outreach sequence — source-linked drafts\n\n")
        for item in lead.get("outreach_sequence", []):
            email_links = " · ".join(
                f"[{sid}]({srcmap[sid]['url']})" for sid in item.get("source_ids", []) if sid in srcmap
            )
            out.append(f"**{item['day']} — {item['subject']}**\n\n{item['body']}\n\n**Evidence to recheck:** {email_links or 'See the source register above.'}\n\n")
        out.append("**Source record:** `data/batch7_records.json`; **contact provenance:** " + lead["contact_source"] + "\n\n")

    out.extend(["## Batch 7 scope decisions\n\n"])
    out.extend(f"- {item}\n" for item in data.get("scope_decisions", []))
    excluded = data.get("excluded_candidates", [])
    if excluded:
        out.extend(["\n### Excluded candidates — not in Batch 7 leads or registry\n\n"])
        for candidate in excluded:
            out.extend([f"#### {candidate['brand']}\n\n**Decision:** {candidate['decision']}\n\n**Reason:** {candidate['reason']}\n\n"])
            excluded_rows = [
                (source["label"], "Candidate source — excluded; not counted as a lead", f"[{source['url']}]({source['url']})", f"Checked {source['checked_on']}. {source['observed']}")
                for source in candidate.get("sources", [])
            ]
            out.append(_md_table(excluded_rows) + "\n\n")
    out.extend([
        "### Batch 7 incremental builder\n\n",
        "```bash\n.venv/bin/python niches/food-bengaluru/scripts/build_batch7_deliverables.py\n```\n\n",
        "The builder appends new source records, contacts, email/source maps and score/finding rows without rebuilding earlier batches; its preflight checks the root registry and every existing niche's lead JSON/CSV inventories. It refuses a second run against existing Batch 7 outputs.\n\n",
        step0_text.rstrip(),
        "\n",
    ])
    return "".join(out)

def append_verify_guide(leads, step0_text):
    existing = VERIFY_MD.read_text(encoding="utf-8")
    markers = [
        "## Batch 7 — Food / Bengaluru (source check: 06 October 2026)",
        STEP0_MARKER,
    ]
    positions = [existing.find(marker) for marker in markers if existing.find(marker) >= 0]
    base = existing[:min(positions)].rstrip() if positions else existing.rstrip()
    VERIFY_MD.write_text(base + "\n\n" + build_verify_batch_section(leads, step0_text).rstrip() + "\n", encoding="utf-8")

def build_all_email_appendix(leads):
    parts = [
        "\n\n---\n\n# Batch 7 — Food / Bengaluru (checked 06 Oct 2026)\n\n",
        f"> {len(leads)} new leads · {len(leads) * 4} messages · Day 1 / 3 / 7 / 14. Working target: ten; shortfall not padded. Every factual opener is mapped to public sources, and published contact routes are not presumed to be media buyers. Re-open source links before sending.\n\n",
    ]
    for lead in leads:
        parts.append(f"## {lead['brand']}\n\n")
        parts.append(f"**Published route:** `{lead['contact']['email']}` · {lead['contact'].get('route_role','')}\n\n")
        if lead.get("contact_source_status"):
            parts.append(f"**Contact verification caveat:** {lead['contact_source_status']}\n\n")
        parts.append(f"**Contact provenance:** {lead['contact_source']}\n\n")
        srcmap = {source["id"]: source for source in lead.get("sources_detail", [])}
        for i, (subject, body) in enumerate(lead["outreach"], 1):
            day = ["Day 1", "Day 3", "Day 7", "Day 14"][i - 1]
            ids = lead.get("email_source_ids", [[]] * 4)[i - 1]
            links = " · ".join(f"[{sid}]({srcmap[sid]['url']})" for sid in ids if sid in srcmap)
            parts.append(f"### Email {i} ({day}) — {subject}\n\n```text\n{body.strip()}\n```\n\n**Evidence to recheck:** {links}\n\n")
    return "".join(parts).rstrip() + "\n"


def build_readme(leads, step0_text, total_registered):
    existing = README_MD.read_text(encoding="utf-8")
    b7_marker = "## Batch 7 — new leads"
    if b7_marker in existing:
        existing = existing[:existing.index(b7_marker)].rstrip()

    replacements = [
        (
            "> **60 leads total:** 21 legacy leads (Batches 1–2) and 39 source-checked leads across Batches 3–6. Batch 6 adds 9 qualified leads, checked 05 Oct 2026 (working target: 10); earlier batches and records are preserved. The 21 legacy leads were not revalidated in this update.",
            f"> **{60 + len(leads)} leads total:** 21 legacy leads (Batches 1–2) and {39 + len(leads)} source-checked leads across Batches 3–7. Batch 7 adds {len(leads)} qualified leads, checked 06 Oct 2026 (working target: 10); earlier batches and records are preserved. The 21 legacy leads were not revalidated in this update.",
        ),
        ("Master workbook with 60 leads, 240 emails", f"Master workbook with {60 + len(leads)} leads, {240 + 4 * len(leads)} emails"),
        (
            "| `Food-Bangalore-BATCH-6-EMAILS.xlsx` | New one-sheet export of the 9 retained Batch 6 leads with Subject, Body 1–4, and audit attachment name. |",
            "| `Food-Bangalore-BATCH-6-EMAILS.xlsx` | Preserved one-sheet export of the 9 retained Batch 6 leads with Subject, Body 1–4, and audit attachment name. |\n| `Food-Bangalore-BATCH-7-EMAILS.xlsx` | New one-sheet export of the six retained Batch 7 leads with Subject, Body 1–4, audit attachment, day-level source IDs, and linked contact provenance/caveats. |",
        ),
        (
            "60 per-lead PDFs. Batch 3, Batch 4, Batch 5 and Batch 6 reports are checked for exactly five pages; previous formats are preserved and not revalidated here.",
            f"{60 + len(leads)} per-lead PDFs. Batch 3–6 five-page reports are preserved; Batch 7 has {len(leads)} source-linked audits with page counts, footer numbering and logo presence checked during generation.",
        ),
        (
            "Claim-to-public-URL rows for prior batches plus Batch 5 and Batch 6 findings, exact ad cards, Google caveats, source register and contact provenance.",
            "Claim-to-public-URL rows for prior batches plus Batch 5–7 findings, exact ad cards, Google caveats, source register and contact provenance.",
        ),
        (
            "240 email messages across the preserved batches; Batch 6 adds 36 new messages.",
            f"{240 + 4 * len(leads)} email messages across the preserved batches; Batch 7 adds {4 * len(leads)} new messages.",
        ),
        (
            "Per-lead, four-touch email files; Batch 5 and Batch 6 routes and contact provenance are explicit.",
            "Per-lead, four-touch email files; Batch 5–7 routes and contact provenance are explicit, including the partial The Kind email caveat.",
        ),
        (
            "60 structured lead records. Batch 6 JSON includes contact provenance, direct Meta IDs, source map, findings, score bases and four outreach emails.",
            f"{60 + len(leads)} structured lead records. Batch 7 JSON includes contact provenance, direct Meta IDs, source map, findings, score bases and four outreach emails.",
        ),
        (
            "Flat index for all 60 leads, with batch vintage and Google status caveats.",
            f"Flat index for all {60 + len(leads)} leads, with batch vintage and Google status caveats.",
        ),
        (
            "`data/batch3_records.json` / `data/batch4_records.json` / `data/batch5_records.json` / `data/batch6_records.json` | Canonical source records for Batches 3–6; Batch 6 stores source notes, findings, score bases and email briefs used by the builder. |",
            "`data/batch3_records.json` / `data/batch4_records.json` / `data/batch5_records.json` / `data/batch6_records.json` / `data/batch7_records.json` | Canonical source records for Batches 3–7; Batch 7 stores source notes, findings, score bases, exclusions and email briefs used by the builder. |",
        ),
        (
            "Updated deduplication registry; 160 registered brands after Batch 6.",
            f"Updated deduplication registry; {total_registered} registered brands after Batch 7.",
        ),
        (
            "| `scripts/build_batch6_deliverables.py` | Batch 6 one-time incremental append builder; its duplicate guard blocks another run against the completed pack. Do not run legacy-only `build_master_excel.py`. |",
            "| `scripts/build_batch6_deliverables.py` | Preserved Batch 6 one-time incremental append builder; do not rerun against its completed pack. |\n| `scripts/build_batch7_deliverables.py` | Batch 7 one-time incremental append builder; preflight blocks duplicate brands/artifacts. Do not run legacy-only `build_master_excel.py`. |",
        ),
        (
            "| Batch 6 | 9 | 05 Oct 2026 | New source-checked batch; one selected evidence-led audit per lead; working target 10, no padding. |",
            "| Batch 6 | 9 | 05 Oct 2026 | Preserved source-checked batch; one selected evidence-led audit per lead. |\n" + f"| Batch 7 | {len(leads)} | 06 Oct 2026 | New source-checked batch; one selected evidence-led audit per lead; working target 10, no padding. |",
        ),
    ]
    for old, new in replacements:
        if old not in existing:
            raise RuntimeError(f"README update could not find expected existing text: {old[:120]}")
        existing = existing.replace(old, new, 1)

    rows = ["", "## Batch 7 — new leads", "",
            "| Brand | Published first-party route | Google evidence status | Public-evidence fit |",
            "|---|---|---|---:|"]
    for lead in leads:
        caveat = " (re-open before sending)" if lead.get("contact_source_status", "").startswith("PARTIALLY VERIFIED") else ""
        rows.append(f"| {lead['brand']} | `{lead['contact']['email']}`{caveat} | {lead['google_evidence']['status']} | {lead['score']} / 10 |")
    rows.extend([
        "", "### Batch 7 evidence and scope notes", "",
        f"- Working target: 10. {len(leads)} leads were retained after individual Meta-card, first-party contact, Bengaluru-context and duplicate checks; no weaker candidate was used to fill the shortfall.",
        f"- All {len(leads)} selected Meta cards were opened individually on 06 Oct 2026 and displayed the named page, Sponsored and Active labels, start date and direct Library ID. One card is not an account-wide count; public inventory can change.",
        "- Duplicate checks include `data/seen_leads.json` and every existing `niches/` folder's lead JSON, CSV indexes and niche-local registries. No duplicate name/domain was found among the six retained brands.",
        "- The Kind: the official-domain indexed `/shop/` Get in Touch snippet showed `thekindindia@gmail.com`, but direct page fetch redirected and did not expose it. Re-open the live official page before sending; the contact is not a confirmed media buyer.",
        "- United Telugu Kitchens' `gm@telugukitchens.com` is published on its Catering page specifically for catering enquiries. La Pino'z lists `support@lapinozpizza.in` for customer support and `info@lapinozpizza.in` for franchise enquiries. Other published inboxes are routing contacts only; no paid-media owner was identified.",
        "- GO DESi's opened Google result names verified agency Haystack Marketing Services and does not expose a GO DESi creative tie. Deliciae's brand-named verified Google creative is removed and last shown 05 Oct; La Pino'z's brand-named creative under verified advertiser COPENHAGEN HOSPITALITY was last shown 05 Oct and is not labelled Active. These are not current-active Google claims. Zero-result domain queries do not prove no ads.",
        "- The Happy Screens was excluded because the official offer and opened ad foreground private-theatre and celebration booking; food is supporting. It is not included as a lead, score, email target or registry entry; the scope decision and source links are in `VERIFY-THE-DATA.md` and the master workbook.",
        "- All prices, product details, offers and quality/volume claims shown in creatives are advertiser copy unless explicitly stated otherwise; they were not independently verified as current terms or outcomes. Spend, ROAS, CPA, order/booking outcomes, tags/pixels, audience geography and event firing remain unavailable/not measured.",
        "- Scores are public-evidence fit only (/20, displayed /10) with five recorded bases; they are not performance, revenue, likelihood, spend, projections, benchmarks, expected uplift or ROI.",
        "", step0_text, "",
        "### Batch 7 builder (one-time append)", "", "```bash",
        ".venv/bin/python niches/food-bengaluru/scripts/build_batch7_deliverables.py", "```", "",
        "The incremental builder appends to the master workbook, email exports, source index, lead/outreach folders, verification guide and registry. It refuses duplicate records and pre-existing Batch 7 output files; the data builder preserves earlier batches and creates no estimated metrics.",
    ])
    return existing.rstrip() + "\n\n" + "\n".join(rows) + "\n"

def _style_new_row(ws, row_idx, template_row=None):
    if template_row is None:
        template_row = max(1, row_idx - 1)
    for col in range(1, ws.max_column + 1):
        src = ws.cell(template_row, col)
        dst = ws.cell(row_idx, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)


def append_xlsx_row(ws, values, hyperlinks=None):
    row_idx = ws.max_row + 1
    template = max(2, row_idx - 1)
    _style_new_row(ws, row_idx, template)
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row_idx, col, value)
        cell.alignment = copy(cell.alignment)
        cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
        if hyperlinks and col in hyperlinks:
            cell.hyperlink = hyperlinks[col]
            cell.style = "Hyperlink"
    return row_idx


def add_vintage_column(ws, brand_col=1, title="Batch / data vintage", min_row=2):
    col = next((c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value == title), ws.max_column + 1)
    header = ws.cell(1, col)
    if not header.value:
        header.value = title
        if col > 1 and ws.cell(1, col - 1).has_style:
            header._style = copy(ws.cell(1, col - 1)._style)
        header.font = copy(header.font)
        header.font = header.font.copy(bold=True, color="FFFFFF")
        header.fill = copy(ws.cell(1, 1).fill)
        header.alignment = header.alignment.copy(wrap_text=True, vertical="center")
    for r in range(min_row, ws.max_row + 1):
        brand = ws.cell(r, brand_col).value
        cell = ws.cell(r, col)
        cell.value = vintage_for(str(brand)) if brand else "General note / cross-pack"
        cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
    from openpyxl.utils import get_column_letter
    ws.column_dimensions[get_column_letter(col)].width = 34
    return col

def add_excel_source_hyperlink(cell, url):
    cell.hyperlink = url
    cell.style = "Hyperlink"


def build_master_workbook(data, leads, page_counts, total_registered, step0_text):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = load_workbook(WORKBOOK)
    required = {
        "README", "Leads", "Emails", "Findings (severity+evidence)", "Opportunity Score",
        "Coverage (measured vs not)", "Evidence Register", "Verify These (links)",
        "Audit Basis (sources)", "Emails (one sheet)", "Legacy ROI Archive", "Batch & Data Status",
    }
    assert required.issubset(set(wb.sheetnames)), f"Workbook structure changed: missing {required - set(wb.sheetnames)}"

    if "Legacy ROI Archive" in wb.sheetnames:
        roi = wb["Legacy ROI Archive"]
        if roi.max_row >= 2 and roi.cell(2, 1).value:
            roi.cell(2, 1).value = "Retained for traceability only. Source-checked Batches 3–7 add no spend, ROAS, CPA, conversion, uplift or ROI projections."

    # Leads — append only; keep all earlier source snapshots and IDs.
    ws = wb["Leads"]
    ws.cell(1, 6).value = "Public-evidence fit /10 (triage)"
    ws.cell(1, 7).value = "Google Ads evidence status"
    ws.cell(1, 13).value = "Audit PDF (page count in Batch & Data Status)"
    if ws.cell(1, 16).value is None:
        ws.cell(1, 16).value = "Batch / data vintage"
        ws.cell(1, 16)._style = copy(ws.cell(1, 15)._style)
    ws.column_dimensions[get_column_letter(16)].width = 34
    lead_detail_headers = [
        "Batch 7 headline", "Public market signal", "Bengaluru evidence status",
        "Published contact route role", "Contact verification status",
        "Meta Library ID / direct card", "Public-evidence score note",
    ]
    for col, header in enumerate(lead_detail_headers, start=17):
        cell = ws.cell(1, col)
        cell.value = header
        cell._style = copy(ws.cell(1, 16)._style)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = [52, 70, 48, 48, 58, 54, 58][col - 17]
    for r in range(2, ws.max_row + 1):
        brand = ws.cell(r, 2).value
        ws.cell(r, 16).value = vintage_for(str(brand)) if brand else ""
        ws.cell(r, 16).alignment = Alignment(wrap_text=True, vertical="top")
    lead_number = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value)
    for lead in leads:
        lead_number += 1
        c = lead["contact"]
        row = append_xlsx_row(ws, [
            lead_number, lead["brand"], lead["website"], lead.get("instagram", ""), lead["niche"], lead["score"],
            lead["ads_active"], c.get("email", ""), c.get("phone", ""), c.get("entity", ""),
            c.get("address", ""), lead["contact_source"],
            f"audits/{lead['slug']}-paid-media-measurement-audit.pdf",
            f"outreach/{lead['slug']}-outreach-templates.md", lead["verified_on"],
            "BATCH 7 — source-checked 06 OCT 2026",
            lead.get("headline", ""), lead.get("market_signal", ""), lead.get("local_evidence_status", ""),
            c.get("route_role", ""), lead.get("contact_source_status", "PUBLISHED FIRST-PARTY ROUTE — not a confirmed paid-media buyer"),
            f"Meta Library ID {lead['ad_evidence'][0]['library_id']} — {lead['ad_evidence'][0]['direct_url']}",
            lead.get("score_note", ""),
        ])
        card_cell = ws.cell(row, 22)
        card_cell.hyperlink = lead["ad_evidence"][0]["direct_url"]
        card_cell.style = "Hyperlink"
    ws.auto_filter.ref = f"A1:W{ws.max_row}"

    # Four individual messages per prospect with their evidence IDs.
    ws = wb["Emails"]
    if ws.cell(1, 7).value is None:
        ws.cell(1, 7).value = "Batch / data vintage"
        ws.cell(1, 7)._style = copy(ws.cell(1, 6)._style)
    if ws.cell(1, 8).value is None:
        ws.cell(1, 8)._style = copy(ws.cell(1, 7)._style)
    ws.cell(1, 8).value = "Email source IDs"
    ws.column_dimensions[get_column_letter(7)].width = 34
    ws.column_dimensions[get_column_letter(8)].width = 22
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        for n, (subject, body) in enumerate(lead["outreach"], start=1):
            day = ["Day 1", "Day 3", "Day 7", "Day 14"][n - 1]
            source_ids = ", ".join(lead["email_source_ids"][n - 1])
            append_xlsx_row(ws, [lead["brand"], n, day, subject, body.strip(), lead["contact"]["email"], "BATCH 7 — checked 06 OCT 2026", source_ids])
    ws.auto_filter.ref = f"A1:J{ws.max_row}"

    # Findings and the corresponding action/verification boundaries.
    ws = wb["Findings (severity+evidence)"]
    if ws.cell(1, 7).value is None:
        ws.cell(1, 7).value = "Next check"
        ws.cell(1, 7)._style = copy(ws.cell(1, 6)._style)
    if ws.cell(1, 8).value is None:
        ws.cell(1, 8).value = "Batch / data vintage"
        ws.cell(1, 8)._style = copy(ws.cell(1, 6)._style)
    ws.column_dimensions[get_column_letter(7)].width = 70
    ws.column_dimensions[get_column_letter(8)].width = 34
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 8).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        for n, finding in enumerate(lead["findings"], start=1):
            append_xlsx_row(ws, [lead["brand"], n, finding.get("priority", "NOTE").upper(), finding["title"], finding["evidence"], finding["analysis"], finding["next_check"], "BATCH 7 — checked 06 OCT 2026"])
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    # Five-dimension public-evidence-fit score with every basis retained.
    ws = wb["Opportunity Score"]
    if ws.cell(1, 8).value is None:
        ws.cell(1, 8).value = "Batch / data vintage"
        ws.cell(1, 8)._style = copy(ws.cell(1, 7)._style)
    ws.column_dimensions[get_column_letter(8)].width = 34
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 8).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        for item in lead["score_breakdown"]:
            append_xlsx_row(ws, [lead["brand"], item["dimension"], item["score"], item["basis"], lead["score_20"], lead["score_band"], lead["score_note"], "BATCH 7 — checked 06 OCT 2026"])
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    # Explicitly measured, unavailable and not-measured surfaces.
    ws = wb["Coverage (measured vs not)"]
    if ws.cell(1, 6).value is None:
        ws.cell(1, 6).value = "Batch / data vintage"
        ws.cell(1, 6)._style = copy(ws.cell(1, 5)._style)
    ws.column_dimensions[get_column_letter(6)].width = 34
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 6).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        for item in lead["coverage"]:
            append_xlsx_row(ws, [lead["brand"], item["surface"], item["status"], item["method"], item["date"], "BATCH 7 — checked 06 OCT 2026"])
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # Evidence Register — one ad-card row plus every dated source record.
    ws = wb["Evidence Register"]
    if ws.cell(1, 6).value is None:
        ws.cell(1, 6).value = "Batch / data vintage"
        ws.cell(1, 6)._style = copy(ws.cell(1, 5)._style)
    ws.column_dimensions[get_column_letter(6)].width = 34
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 6).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        srcmap = {source["id"]: source for source in lead["sources_detail"]}
        for ad in lead.get("ad_evidence", []):
            src = srcmap[ad["source_id"]]
            observed = f"{ad['status']}; started {ad.get('started_on','not stated')}; page {ad.get('page_name','')}; page ID {ad.get('page_id','')}; copy: {ad.get('creative_excerpt','')}; CTA: {ad.get('cta','')}; destination: {ad.get('destination_status','')}"
            append_xlsx_row(ws, [lead["brand"], f"Meta Library ID {ad['library_id']}", src["url"], observed, "06 OCT 2026", "BATCH 7 — checked 06 OCT 2026"])
            add_excel_source_hyperlink(ws.cell(ws.max_row, 3), src["url"])
        for source in lead.get("sources_detail", []):
            append_xlsx_row(ws, [lead["brand"], f"{source['id']} — {source['label']}", source["url"], source.get("observed", ""), source.get("verified_on", "2026-10-06"), "BATCH 7 — checked 06 OCT 2026"])
            add_excel_source_hyperlink(ws.cell(ws.max_row, 3), source["url"])
    for candidate in data.get("excluded_candidates", []):
        for source in candidate.get("sources", []):
            append_xlsx_row(ws, [candidate["brand"], f"EXCLUDED — {source['label']}", source["url"], source.get("observed", ""), source.get("checked_on", "2026-10-06"), "BATCH 7 — excluded candidate; checked 06 OCT 2026"])
            add_excel_source_hyperlink(ws.cell(ws.max_row, 3), source["url"])
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # Audit Basis — retain a linked row for each first-party or platform source.
    ws = wb["Audit Basis (sources)"]
    if ws.cell(1, 3).value is None:
        ws.cell(1, 3).value = "Batch / data vintage"
        ws.cell(1, 3)._style = copy(ws.cell(1, 2)._style)
    ws.column_dimensions[get_column_letter(3)].width = 34
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 3).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        for source in lead.get("sources_detail", []):
            text = f"{source['id']} — {source['label']}. Observed: {source.get('observed','')} URL: {source['url']}"
            row = append_xlsx_row(ws, [lead["brand"], text, "BATCH 7 — checked 06 OCT 2026"])
            ws.cell(row, 2).hyperlink = source["url"]
            ws.cell(row, 2).style = "Hyperlink"
    for candidate in data.get("excluded_candidates", []):
        for source in candidate.get("sources", []):
            text = f"EXCLUDED — {candidate['brand']}: {candidate['decision']} {candidate['reason']} Source: {source['label']} {source['observed']} URL: {source['url']}"
            row = append_xlsx_row(ws, [candidate["brand"], text, "BATCH 7 — excluded candidate; checked 06 OCT 2026"])
            ws.cell(row, 2).hyperlink = source["url"]
            ws.cell(row, 2).style = "Hyperlink"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"

    # Claim-level source links plus contact and exclusion provenance.
    ws = wb["Verify These (links)"]
    if ws.cell(1, 6).value is None:
        ws.cell(1, 6).value = "Batch / data vintage"
        ws.cell(1, 6)._style = copy(ws.cell(1, 5)._style)
    ws.column_dimensions[get_column_letter(6)].width = 34
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 6).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        srcmap = {source["id"]: source for source in lead.get("sources_detail", [])}
        for source in lead.get("sources_detail", []):
            append_xlsx_row(ws, [lead["brand"], f"{source['id']} — {source['label']}", source["source_type"], source["url"], source.get("observed", ""), "BATCH 7 — checked 06 OCT 2026"], hyperlinks={4: source["url"]})
        for finding in lead.get("findings", []):
            for sid in finding.get("source_ids", []):
                source = srcmap[sid]
                append_xlsx_row(ws, [lead["brand"], finding["title"], f"Supporting finding source {sid}", source["url"], f"{finding['evidence']} Analysis: {finding['analysis']} Next check: {finding['next_check']}", "BATCH 7 — checked 06 OCT 2026"], hyperlinks={4: source["url"]})
        for n, item in enumerate(lead.get("outreach_sequence", []), start=1):
            for sid in item.get("source_ids", []):
                source = srcmap[sid]
                append_xlsx_row(ws, [lead["brand"], f"{item['day']} email source {sid}", "Email evidence map", source["url"], f"Subject: {item['subject']}\nEmail body retained in Emails (one sheet) and Emails sheets.", "BATCH 7 — checked 06 OCT 2026"], hyperlinks={4: source["url"]})
    for candidate in data.get("excluded_candidates", []):
        for source in candidate.get("sources", []):
            append_xlsx_row(ws, [candidate["brand"], "Scope exclusion — primary offer not core Food", "Excluded candidate source", source["url"], f"{candidate['decision']} {candidate['reason']} Evidence: {source['observed']}", "BATCH 7 — excluded candidate; checked 06 OCT 2026"], hyperlinks={4: source["url"]})
    for decision in data.get("scope_decisions", []):
        append_xlsx_row(ws, ["Batch 7 scope", "Scope decision", "Scope note / source map in VERIFY-THE-DATA.md", "", decision, "BATCH 7 — checked 06 OCT 2026"])
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # One row per lead with the full four-email sequence and source IDs by day.
    ws = wb["Emails (one sheet)"]
    if ws.cell(1, 9).value is None:
        ws.cell(1, 9).value = "Batch / data vintage"
        ws.cell(1, 9)._style = copy(ws.cell(1, 8)._style)
    if ws.cell(1, 10).value is None:
        ws.cell(1, 10)._style = copy(ws.cell(1, 9)._style)
    ws.cell(1, 10).value = "Email source IDs by day"
    ws.column_dimensions[get_column_letter(9)].width = 34
    ws.column_dimensions[get_column_letter(10)].width = 36
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 9).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 9).alignment = Alignment(wrap_text=True, vertical="top")
    for lead in leads:
        subjects = "\n".join(f"{day}: {item[0]}" for day, item in zip(["Day 1", "Day 3", "Day 7", "Day 14"], lead["outreach"]))
        bodies = [item[1].strip() for item in lead["outreach"]]
        source_map = "; ".join(f"{day}: {', '.join(ids)}" for day, ids in zip(["Day 1", "Day 3", "Day 7", "Day 14"], lead["email_source_ids"]))
        append_xlsx_row(ws, [lead["brand"], lead["contact"]["email"], subjects, *bodies, f"{lead['slug']}-paid-media-measurement-audit.pdf", "BATCH 7 — checked 06 OCT 2026", source_map])
    ws.auto_filter.ref = f"A1:J{ws.max_row}"

    # Refresh batch labels only in the two mixed-batch legacy reference sheets.
    for title in ("Leaks & Roadmap", "You vs Competitor"):
        if title in wb.sheetnames:
            add_vintage_column(wb[title], brand_col=1)

    # Record Step 0 dispositions without changing prior batch rows.
    ws = wb["Skills Used"]
    skill_rows = [
        ["Batch 7 public-source research", "vamsy16/arena-performance-marketing: REUSABLE-PROMPT.md and Food/Bengaluru pack; vamsy16/claude-ads @ 669c7608ecb50dd95c941a71fa3ca0a1c0e40512: skills/ads-research", "Used the current archive and source/date lineage. Individual Meta cards and brand-owned contact/local pages were opened; keyword-only results were not treated as advertiser proof."],
        ["Batch 7 evidence coverage and measurement limits", "vamsy16/claude-ads @ 669c7608ecb50dd95c941a71fa3ca0a1c0e40512: skills/ads-audit; vamsy16/marketingskills @ 59d5112e61fd9b043cb3c97d5551f7044d184ad0: skills/analytics, skills/attribution, skills/ads/references/audit-guardrails.md", "Used evidence coverage, unknown/partial boundaries and event-definition/attribution limits. No authenticated account audit, performance score or pixel-absence claim is represented."],
        ["Batch 7 outreach", "vamsy16/arena-performance-marketing: skills/cold-email and references; vamsy16/smart-pursuit-agency @ 40e070ceaaf632790a03a644e6d3eb7a5dce40b2: lead-generation-prospecting.md and ppc-paid-media.md", "Used one verifiable public observation per touch and a low-friction routing ask. Four source-linked Day 1/3/7/14 drafts; no messages sent and no unsupported metrics copied."],
        ["Batch 7 pack and PDF format", "vamsy16/arena-performance-marketing: REUSABLE-PROMPT.md; niches/food-bengaluru/scripts/build_batch6_deliverables.py schema/layout adapted into build_batch7_deliverables.py", "Used the existing branded public-evidence audit and incremental append pattern. PDF pages/footer numbers are measured after render; no ROI/guarantee format was used."],
        ["Batch 7 excluded workflow examples", "vamsy16/arena-performance-marketing: AUDIT-WORKFLOW.md and HOW-TO-RUN-IN-NEW-CHAT.md; Step 0 review in VERIFY-THE-DATA.md", "Mock/generated lead pools, Maps-scraping workflows and estimated performance examples were not run or copied. Repositories with no relevant usable skill and unrelated repositories are recorded in the Step 0 inventory."],
    ]
    for row in skill_rows:
        append_xlsx_row(ws, row)

    # Batch & Data Status — selected leads plus a separate non-lead scope exclusion.
    status = wb["Batch & Data Status"]
    seen_status = {norm_name(str(status.cell(r, 1).value or "")) for r in range(2, status.max_row + 1)}
    for lead in leads:
        if norm_name(lead["brand"]) in seen_status:
            raise RuntimeError(f"{lead['brand']} already appears in Batch & Data Status")
        pages = page_counts[lead["brand"]]
        row = append_xlsx_row(status, [lead["brand"], "Batch 7", "06 OCT 2026", "NEW — source-checked 06 OCT 2026",
                                      f"{pages}-page source-linked public-evidence audit; public-evidence fit /20 and /10",
                                      "Dated public snapshot; private performance and conversion outcomes not measured."])
        status.cell(row, 4).fill = PatternFill("solid", fgColor="D1FAE5")
        seen_status.add(norm_name(lead["brand"]))
    for candidate in data.get("excluded_candidates", []):
        if norm_name(candidate["brand"]) in seen_status:
            continue
        row = append_xlsx_row(status, [candidate["brand"], "Batch 7 — excluded, not a lead", "06 OCT 2026", "EXCLUDED — scope decision",
                                      "Not scored; primary marketed service is private theatre/celebrations, with food a supporting feature.",
                                      "Source links and reason recorded in VERIFY-THE-DATA.md and Evidence Register."])
        status.cell(row, 4).fill = PatternFill("solid", fgColor="FEF3C7")
    status.freeze_panes = "A2"
    status.auto_filter.ref = status.dimensions
    status.sheet_view.showGridLines = False

    # Workbook landing page: append counts accurately, retain all earlier batch history.
    ws = wb["README"]
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    count = len(leads)
    total_leads = 60 + count
    total_emails = 240 + count * 4
    pages_summary = ", ".join(f"{brand}: {pages}p" for brand, pages in page_counts.items())
    readme_rows = [
        ["SMART PURSUIT — FOOD / BENGALURU EVIDENCE PACK", ""],
        [f"{total_leads} leads · {total_leads} reports · {total_emails} email messages. Batch 7 adds {count} source-checked leads; working target 10, no padding; Batches 1–6 are preserved.", ""],
        ["DATA VINTAGE", ""],
        ["Batch 1 + Batch 2", "21 earlier leads captured 02–03 OCT 2026; retained as legacy snapshots and not revalidated in this update."],
        ["Batch 3", "10 leads checked 04 OCT 2026; prior five-page evidence-focused audits and outreach are preserved."],
        ["Batch 4", "10 leads checked 04 OCT 2026; prior five-page evidence-focused audits and outreach are preserved."],
        ["Batch 5", "10 leads checked 04 OCT 2026; prior five-page evidence-focused audits and outreach are preserved."],
        ["Batch 6", "9 leads checked 05 OCT 2026; source-linked audits and four-message sequences are preserved."],
        ["Batch 7", f"{count} leads checked 06 OCT 2026; {count} source-linked audits, {count * 4} emails, source maps, contact provenance, findings and five-dimension public-evidence-fit scores. Per-lead PDF page counts: {pages_summary}. Working target 10; no padding."],
        ["Score meaning", "Public-evidence fit only (five dimensions, /20 displayed /10); not performance, ROI, sales probability, forecast, spend or expected uplift."],
        ["Legacy ROI content", "The Legacy ROI Archive is retained for traceability only; it is not a current forecast. No Batches 3–7 projections are added."],
        ["Source integrity", "Contact routes come from brand-owned pages or official-domain indexed text. The Kind's email remains partially verified and must be re-opened before sending; UTK's email is catering-specific. The Happy Screens is excluded for theatre/celebration-first scope."],
        ["Limits", "Spend, ROAS, CPA, orders, bookings, conversions, site tags/pixels, event firing and audience geography are unavailable/not measured. No estimate, benchmark, uplift or projected result is added."],
        ["Google scope", "Domain-query counts are dated any-time snapshots. GO DESi's opened agency creative is not tied to the brand; Deliciae's creative is removed; La Pino'z last-shown archive item is not labelled Active. Zero results are not proof of no ads."],
        ["Meta scope", f"{count} selected individual Meta cards displayed named page, Sponsored and Active labels, start date and ID on 06 OCT 2026; one selected card is not an account-wide count and public inventory changes daily."],
        ["Local scope", "City relevance comes from explicit selected creative or first-party material; ad audience geography is not inferred."],
        ["Scope decision", "Six leads passed name/domain duplicate checks across the root registry and existing niche JSON/CSV inventories. The Happy Screens was excluded rather than used to pad the working target."],
        ["Workbook coverage", "Leads · Emails · Findings · Opportunity Score · Coverage · Evidence Register · Verify These · Audit Basis · email exports · Skills Used · Batch & Data Status."],
        ["Skills", "Detailed repository/skill dispositions and workflow exclusions are recorded in Step 0 in VERIFY-THE-DATA.md and README.md."],
        ["Build", "Batch 7 used niches/food-bengaluru/scripts/build_batch7_deliverables.py for an incremental append. Do not rerun it against the completed registry/workbook or run legacy-only build_master_excel.py."],
    ]
    for r, row in enumerate(readme_rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    for r in range(2, len(readme_rows) + 1):
        ws.cell(r, 1).font = Font(bold=True, color=TEAL if r == 3 else NAVY)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(1, len(readme_rows) + 1):
        ws.row_dimensions[r].height = 34 if r > 2 else 28
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 108
    ws.sheet_view.showGridLines = False

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(WORKBOOK)


def update_email_exports(leads):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # Append to the existing standalone email workbook, preserving old messages.
    wb = load_workbook(EMAIL_WORKBOOK)
    ws = wb["Emails"]
    if ws.cell(1, 9).value is None:
        ws.cell(1, 9).value = "Batch / data vintage"
        ws.cell(1, 9)._style = copy(ws.cell(1, 8)._style)
        ws.column_dimensions[get_column_letter(9)].width = 34
    if ws.cell(1, 10).value is None:
        ws.cell(1, 10)._style = copy(ws.cell(1, 9)._style)
    ws.cell(1, 10).value = "Email source IDs by day"
    ws.column_dimensions[get_column_letter(10)].width = 36
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 9).value = vintage_for(str(ws.cell(r, 1).value))
        ws.cell(r, 9).alignment = Alignment(wrap_text=True, vertical="top")
    for l in leads:
        subjects = "\n".join(f"{day}: {item[0]}" for day, item in zip(["Day 1", "Day 3", "Day 7", "Day 14"], l["outreach"]))
        bodies = [item[1].strip() for item in l["outreach"]]
        source_map = "; ".join(f"{day}: {', '.join(ids)}" for day, ids in zip(["Day 1", "Day 3", "Day 7", "Day 14"], l["email_source_ids"]))
        append_xlsx_row(ws, [l["brand"], l["contact"]["email"], subjects, *bodies, f"{l['slug']}-paid-media-measurement-audit.pdf", "BATCH 7 — checked 06 OCT 2026", source_map])
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    wb.save(EMAIL_WORKBOOK)

    # CSV is a flat export of the same one-row-per-brand sequence sheet.
    all_rows = []
    for l in leads:
        subjects = "\n".join(f"{day}: {item[0]}" for day, item in zip(["Day 1", "Day 3", "Day 7", "Day 14"], l["outreach"]))
        bodies = [item[1].strip() for item in l["outreach"]]
        source_map = "; ".join(f"{day}: {', '.join(ids)}" for day, ids in zip(["Day 1", "Day 3", "Day 7", "Day 14"], l["email_source_ids"]))
        all_rows.append([l["brand"], l["contact"]["email"], subjects, *bodies, f"{l['slug']}-paid-media-measurement-audit.pdf", "BATCH 7 — checked 06 OCT 2026", source_map])
    with EMAIL_CSV.open("r", newline="", encoding="utf-8-sig") as f:
        old = list(csv.reader(f))
    if old and any("BATCH 7 — checked 06 OCT 2026" in row for row in old[1:]):
        raise RuntimeError("Batch 7 email CSV rows already exist; refusing duplicate append")
    if old and len(old[0]) == 8:
        old[0].append("Batch / data vintage")
        for row in old[1:]:
            if len(row) == 8:
                row.append(vintage_for(row[0]))
    if old and len(old[0]) == 9:
        old[0].append("Email source IDs by day")
    elif old and len(old[0]) >= 10:
        old[0][9] = "Email source IDs by day"
    for row in old[1:]:
        if len(row) < 10:
            row.extend([""] * (10 - len(row)))
    with EMAIL_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerows(old)
        writer.writerows(all_rows)


def build_batch7_email_workbook(leads):
    """Write the Batch 7 email export with source IDs and contact provenance."""
    import math
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Batch 7 Email Sequences"
    headers = [
        "Brand name", "Email", "Subject", "Body 1", "Body 2", "Body 3", "Body 4", "Attachment Name",
        "Evidence source IDs by day", "Contact source / verification caveat",
    ]
    ws.append(headers)
    days = ["Day 1", "Day 3", "Day 7", "Day 14"]
    for lead in leads:
        subjects = "\n".join(f"{day}: {item[0]}" for day, item in zip(days, lead["outreach"]))
        bodies = [str(item[1]).replace("\r\n", "\n").strip() for item in lead["outreach"]]
        source_map = "; ".join(f"{day}: {', '.join(ids)}" for day, ids in zip(days, lead["email_source_ids"]))
        contact_status = lead.get("contact_source_status", "PUBLISHED FIRST-PARTY ROUTE — not a confirmed paid-media buyer")
        contact_provenance = contact_status + "\n" + lead["contact_source"]
        ws.append([
            lead["brand"], lead["contact"]["email"], subjects, *bodies,
            f"{lead['slug']}-paid-media-measurement-audit.pdf", source_map, contact_provenance,
        ])
        contact_source = next((src for src in lead.get("sources_detail", []) if src["id"] in lead["contact_source_detail"]["source_ids"]), None)
        if contact_source:
            ws.cell(ws.max_row, 10).hyperlink = contact_source["url"]
            ws.cell(ws.max_row, 10).style = "Hyperlink"

    widths = {"A": 25, "B": 36, "C": 48, "D": 64, "E": 64, "F": 64, "G": 64, "H": 52, "I": 42, "J": 76}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    navy_fill = PatternFill("solid", fgColor=NAVY)
    white_bold = Font(name="Aptos", size=10, bold=True, color=WHITE)
    body_font = Font(name="Aptos", size=10, color=INK)
    brand_font = Font(name="Aptos", size=10, bold=True, color=TEAL)
    edge = Side(style="hair", color=RULE)
    cell_border = Border(bottom=edge, right=edge)
    for cell in ws[1]:
        cell.fill = navy_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[1].height = 30
    for row_index in range(2, ws.max_row + 1):
        max_lines = 1
        for col_index in range(1, 11):
            cell = ws.cell(row_index, col_index)
            cell.font = brand_font if col_index == 1 else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = cell_border
            text = str(cell.value or "")
            width = widths[ws.cell(1, col_index).column_letter]
            chars_per_line = max(24, int(width * 1.08))
            visual_lines = sum(
                max(1, math.ceil(len(line) / chars_per_line))
                for line in text.replace("\r\n", "\n").split("\n")
            )
            max_lines = max(max_lines, visual_lines)
        ws.row_dimensions[row_index].height = min(420, max(150, max_lines * 12 + 18))

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = TEAL
    table = Table(displayName="Batch7EmailSequence", ref=f"A1:J{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    wb.properties.title = "Food / Bengaluru — Batch 7 Email Sequences"
    wb.properties.subject = f"{len(leads)} source-checked Batch 7 leads with four outreach emails each"
    wb.properties.creator = "Smart Pursuit"
    if BATCH7_EMAIL_WORKBOOK.exists():
        raise RuntimeError(f"Batch 7 email workbook already exists: {BATCH7_EMAIL_WORKBOOK}")
    wb.save(BATCH7_EMAIL_WORKBOOK)


def update_leads_csv(leads):
    rows = []
    with LEADS_CSV.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)
    proposed_names = {norm_name(l["brand"]) for l in leads}
    if any(norm_name(r.get("brand", "")) in proposed_names for r in rows):
        raise RuntimeError("Batch 7 brand already exists in LEADS-INDEX.csv")
    for name in ("batch", "data_vintage"):
        if name not in fieldnames:
            fieldnames.append(name)
    batch1 = {norm_name(x) for x in LEGACY_BATCH1}
    batch2 = {norm_name(x) for x in LEGACY_BATCH2}
    batch3 = {norm_name(x) for x in BATCH3_LEADS}
    batch4 = {norm_name(x) for x in BATCH4_LEADS}
    batch5 = {norm_name(x) for x in BATCH5_LEADS}
    batch6 = {norm_name(x) for x in BATCH6_LEADS}
    batch7 = {norm_name(x) for x in BATCH7_LEADS}
    for row in rows:
        key = norm_name(row.get("brand", ""))
        if key in batch1:
            row["batch"] = "Batch 1"
        elif key in batch2:
            row["batch"] = "Batch 2"
        elif key in batch3:
            row["batch"] = "Batch 3"
        elif key in batch4:
            row["batch"] = "Batch 4"
        elif key in batch5:
            row["batch"] = "Batch 5"
        elif key in batch6:
            row["batch"] = "Batch 6"
        elif key in batch7:
            row["batch"] = "Batch 7"
        else:
            row["batch"] = row.get("batch") or "Unclassified legacy"
        row["data_vintage"] = vintage_for(row.get("brand", ""))
    for l in leads:
        rows.append({
            "brand": l["brand"], "website": l["website"], "niche": l["niche"],
            "score": l["score"], "google_ads_live": l["ads_active"],
            "lead_email": l["contact"]["email"], "lead_phone": l["contact"].get("phone", ""),
            "contact_entity": l["contact"].get("entity", ""),
            "audit_pdf": f"audits/{l['slug']}-paid-media-measurement-audit.pdf",
            "outreach_md": f"outreach/{l['slug']}-outreach-templates.md",
            "verified_on": l["verified_on"], "batch": "Batch 7",
            "data_vintage": "BATCH 7 — source-checked 06 OCT 2026",
        })
    with LEADS_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

def reconcile_registry(leads, registry):
    brands = registry.setdefault("brands", {})
    new_names = []
    existing_days = [int(row.get("day", -1)) for row in registry.get("daily_log", [])]
    next_day = max(existing_days, default=0) + 1
    if next_day != 9:
        raise RuntimeError(f"Expected to append the next Food/Bengaluru day after Day 8; registry next day is {next_day}")
    for lead in leads:
        key = lead["brand"]
        if key in brands:
            raise RuntimeError(f"Registry already contains {key}")
        brands[key] = {
            "first_seen": "2026-10-06T00:00:00",
            "day": next_day,
            "website": lead["website"],
            "niche": f"Food/Bengaluru/{lead['niche']}",
        }
        new_names.append(key)
    daily = registry.setdefault("daily_log", [])
    if any(int(row.get("day", -1)) == next_day for row in daily):
        raise RuntimeError(f"Root registry already contains Day {next_day}; refusing a second append")
    daily.append({
        "day": next_day, "date": "2026-10-06", "niche": "Food/Bengaluru",
        "count": len(new_names), "total_so_far": len(brands), "brands": new_names,
        "note": "Batch 7 individually source-checked public-source shortlist; earlier batches preserved.",
    })
    registry["total_found"] = len(brands)
    REGISTRY.write_text(json.dumps(registry, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def validate_emails_and_sources(leads):
    for lead in leads:
        assert len(lead["outreach"]) == 4, lead["brand"]
        sources = {source["id"]: source for source in lead.get("sources_detail", [])}
        assert sources and "S1" in sources and "S2" in sources, lead["brand"]
        email = lead["contact"]["email"].casefold()
        source_text = " ".join(
            [lead.get("contact_source", "")] +
            [source.get("observed", "") + " " + source.get("url", "") for source in sources.values()]
        ).casefold()
        assert email and email in source_text, f"Published contact email not present in first-party provenance: {lead['brand']}"
        assert lead["contact_source"] and lead["contact_source_detail"].get("source_ids"), lead["brand"]
        assert set(lead["contact_source_detail"]["source_ids"]) <= set(sources), lead["brand"]
        assert all(set(f.get("source_ids", [])) <= set(sources) for f in lead.get("findings", [])), lead["brand"]
        assert len(lead.get("findings", [])) >= 1, lead["brand"]
        assert len(lead.get("ad_evidence", [])) == 1, lead["brand"]
        sequence = lead.get("outreach_sequence", [])
        assert [item.get("day") for item in sequence] == ["Day 1", "Day 3", "Day 7", "Day 14"], lead["brand"]
        assert all(set(item.get("source_ids", [])) <= set(sources) and item.get("body") and item.get("subject") for item in sequence), lead["brand"]
        ad = lead["ad_evidence"][0]
        assert ad.get("active") is True and ad.get("sponsored") is True, lead["brand"]
        assert ad.get("checked_on") == "2026-10-06" and ad.get("started_on"), lead["brand"]
        assert ad.get("library_id") in ad.get("direct_url", "") and ad.get("page_id") in ad.get("direct_url", ""), lead["brand"]
        assert all(source.get("verified_on") == "2026-10-06" for source in sources.values()), lead["brand"]
        if lead.get("contact_source_status", "").startswith("PARTIALLY VERIFIED"):
            assert "re-open" in lead["contact_source"].casefold() or "re-open" in lead["contact_source_status"].casefold(), lead["brand"]
        for subject, body in lead["outreach"]:
            assert "smartpursuit3@gmail.com" in body and "7095024220" in body, lead["brand"]
            text = (subject + " " + body).casefold()
            for forbidden in ("expected roi", "guaranteed roas", "increase roas by", "reduce cpa by", "conversion rate will", "expected uplift"):
                assert forbidden not in text, f"Unpermitted claim in {lead['brand']}: {forbidden}"


def preflight_outputs(leads, registry):
    expected_paths = []
    for lead in leads:
        expected_paths.extend([
            ROOT / "leads" / f"{lead['slug']}.json",
            ROOT / "outreach" / f"{lead['slug']}-outreach-templates.md",
            ROOT / "audits" / f"{lead['slug']}-paid-media-measurement-audit.pdf",
        ])
    expected_paths.append(BATCH7_EMAIL_WORKBOOK)
    existing = [str(path.relative_to(REPO)) for path in expected_paths if path.exists()]
    if existing:
        raise RuntimeError("Batch 7 output already exists; refusing overwrite: " + ", ".join(existing))

    for lead in leads:
        if any(norm_name(key) == norm_name(lead["brand"]) for key in registry.get("brands", {})):
            raise RuntimeError(f"Registry already contains {lead['brand']}")
    daily = registry.get("daily_log", [])
    if any(int(row.get("day", -1)) == 9 for row in daily):
        raise RuntimeError("Root registry already has Day 9; refusing to append a second batch")

    # Check the consolidated workbook before any output file is written.
    from openpyxl import load_workbook
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        leads_ws = wb["Leads"]
        existing_names = {norm_name(str(row[1] or "")) for row in leads_ws.iter_rows(min_row=2, values_only=True) if len(row) > 1}
        collisions = [lead["brand"] for lead in leads if norm_name(lead["brand"]) in existing_names]
        if collisions:
            raise RuntimeError("Batch 7 brand already exists in master workbook: " + ", ".join(collisions))
    finally:
        wb.close()

    # Consolidated text/CSV exports are separately checked to keep an interrupted rerun safe.
    csv_text = EMAIL_CSV.read_text(encoding="utf-8-sig")
    if "BATCH 7 — checked 06 OCT 2026" in csv_text:
        raise RuntimeError("Batch 7 email CSV rows already exist; refusing duplicate append")
    for path, marker in [
        (ALL_EMAILS_MD, "# Batch 7 — Food / Bengaluru (checked 06 Oct 2026)"),
        (VERIFY_MD, "## Batch 7 — Food / Bengaluru (source check: 06 October 2026)"),
        (README_MD, "## Batch 7 — new leads"),
    ]:
        if marker in path.read_text(encoding="utf-8"):
            raise RuntimeError(f"Batch 7 section already exists in {path.name}; refusing duplicate append")


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Incrementally build Batch 7 Food/Bengaluru deliverables")
    parser.add_argument("--check-only", action="store_true", help="Validate records, contact/source maps and duplicates without writing deliverables")
    args = parser.parse_args()

    data = load_data()
    leads = data["leads"]
    registry = check_deduplication(leads)
    validate_emails_and_sources(leads)
    step0_text = extract_step0_b7(VERIFY_MD.read_text(encoding="utf-8"))
    if args.check_only:
        print(f"PASS — {len(leads)} Batch 7 leads; {len(leads) * 4} source-mapped emails; names/domains are clear in the registry and existing niche JSON/CSV inventories.")
        return

    preflight_outputs(leads, registry)
    page_counts = generate_pdfs(leads)
    write_json_and_outreach(leads)
    append_verify_guide(leads, step0_text)

    all_email_text = ALL_EMAILS_MD.read_text(encoding="utf-8")
    email_marker = "\n\n---\n\n# Batch 7 — Food / Bengaluru (checked 06 Oct 2026)"
    if email_marker in all_email_text:
        all_email_text = all_email_text[:all_email_text.index(email_marker)].rstrip()
    ALL_EMAILS_MD.write_text(all_email_text + build_all_email_appendix(leads), encoding="utf-8")

    total_registered = len(registry.get("brands", {})) + len(leads)
    build_master_workbook(data, leads, page_counts, total_registered, step0_text)
    update_email_exports(leads)
    build_batch7_email_workbook(leads)
    update_leads_csv(leads)
    README_MD.write_text(build_readme(leads, step0_text, total_registered), encoding="utf-8")
    reconcile_registry(leads, registry)

    print(f"Rendered {len(leads)} leads, {len(leads) * 4} emails, {len(leads)} PDFs.")
    print(f"Audit page counts: {page_counts}; total audit pages: {sum(page_counts.values())}.")
    print(f"Updated master workbook: {WORKBOOK}")
    print(f"Updated registry: {len(registry['brands'])} unique entries.")


if __name__ == "__main__":
    main()
